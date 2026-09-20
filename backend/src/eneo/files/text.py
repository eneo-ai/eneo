import asyncio
import logging
import multiprocessing
import signal
import time
import zipfile
from collections.abc import Generator
from contextlib import contextmanager
from dataclasses import dataclass
from enum import Enum
from multiprocessing.process import BaseProcess
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Final, Literal

import magic
import pdfplumber
import pptx
from docx2python import docx2python
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfminer.pdfparser import PDFSyntaxError
from pptx.exc import PackageNotFoundError
from pydantic import BaseModel, TypeAdapter

logger = logging.getLogger(__name__)


# =============================================================================
# Custom Exceptions
# =============================================================================


class ExtractionError(Exception):
    """Base exception for text extraction failures."""

    def __init__(self, message: str, code: str = "EXTRACTION_FAILED"):
        self.message = message
        self.code = code
        super().__init__(self.message)


PdfExtractionLimit = Literal["pages", "extracted_bytes", "seconds"]


@dataclass(frozen=True, slots=True)
class PdfExtractionLimits:
    max_pages: int
    max_extracted_bytes: int
    timeout_seconds: int


class PdfExtractionLimitExceeded(ExtractionError):
    """Raised when PDF text extraction exceeds a deployment ceiling."""

    def __init__(self, limit: PdfExtractionLimit, measured: int | float, ceiling: int):
        self.limit: PdfExtractionLimit = limit
        self.measured = measured
        self.ceiling = ceiling
        super().__init__(
            f"PDF extraction exceeds {limit} limit ({measured} > {ceiling})"
        )


class _PdfLimitResult(BaseModel):
    limit: PdfExtractionLimit
    measured: int | float
    ceiling: int


class _PdfErrorResult(BaseModel):
    kind: Literal["encrypted", "corrupt", "error"]
    details: str


_PDF_RESULT: TypeAdapter[str | _PdfLimitResult | _PdfErrorResult] = TypeAdapter(
    str | _PdfLimitResult | _PdfErrorResult
)


class NoExtractableTextError(ExtractionError):
    """Raised when extraction or transcription produces no usable text."""

    def __init__(self, filename: str):
        super().__init__(
            f"File '{filename}' contains no extractable text", "NO_EXTRACTABLE_TEXT"
        )


class EncryptedFileError(ExtractionError):
    """Raised for password-protected files."""

    def __init__(self, filename: str):
        super().__init__(
            f"File '{filename}' is encrypted/password-protected", "ENCRYPTED"
        )


class CorruptFileError(ExtractionError):
    """Raised for corrupted or malformed files."""

    def __init__(self, filename: str, details: str = ""):
        message = f"File '{filename}' is corrupted or malformed"
        if details:
            message = f"{message}. {details}"
        super().__init__(message, "CORRUPT")


class UnsupportedFormatError(ExtractionError):
    """Raised for unsupported file formats."""

    def __init__(self, filename: str, format_type: str):
        super().__init__(
            f"Format '{format_type}' not supported. "
            f"Please convert '{filename}' to a supported format.",
            "UNSUPPORTED_FORMAT",
        )


# =============================================================================
# MIME Types
# =============================================================================


class MimeTypesBase(str, Enum):
    @classmethod
    def has_value(cls, value: str) -> bool:
        base_value = value.split(";")[0].strip()
        return any(base_value == item.value for item in cls)

    @classmethod
    def values(cls) -> list[str]:
        return [item.value for item in cls]


class TextMimeTypes(MimeTypesBase):
    # Text formats the extractor can handle. Doubles as the upload/attachment
    # allowlist (see limits.limit_service). The crawler maintains its own
    # download policy in crawler.parse_html.CRAWLABLE_DOCUMENT_MIMETYPES — keep
    # the two concerns separate so changing one never silently affects the other.
    # Supported formats
    MD = "text/markdown"
    TXT = "text/plain"
    PDF = "application/pdf"
    DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    TEXT_CSV = "text/csv"
    APP_CSV = "application/csv"
    PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    XLS = "application/vnd.ms-excel"
    JSON = "application/json"
    # Browsers and libmagic disagree on the XML mimetype, so accept both spellings
    XML = "text/xml"
    XML_APP = "application/xml"

    # Legacy formats (for detection/rejection only)
    DOC = "application/msword"
    PPT = "application/vnd.ms-powerpoint"


# =============================================================================
# Text Processing
# =============================================================================

TextExtractionWarning = Literal["pdf_text_likely_reversed"]
PDF_TEXT_LIKELY_REVERSED_WARNING: Final[TextExtractionWarning] = (
    "pdf_text_likely_reversed"
)
TEXT_EXTRACTION_WARNINGS: Final[frozenset[TextExtractionWarning]] = frozenset(
    {PDF_TEXT_LIKELY_REVERSED_WARNING}
)
_PDF_TEXT_MIN_TOKENS_FOR_REVERSAL_WARNING: Final = 40
_PDF_TEXT_MIN_REVERSED_COMMON_TOKENS: Final = 5
_PDF_TEXT_MIN_VERTICAL_LETTER_RUN: Final = 8
_PDF_TEXT_COMMON_WORDS: Final = frozenset(
    (
        "och",
        "att",
        "det",
        "den",
        "som",
        "med",
        "till",
        "eller",
        "ska",
        "har",
        "kan",
        "inte",
    )
)
_PDF_TEXT_REVERSED_COMMON_WORDS: Final = frozenset(
    word[::-1] for word in _PDF_TEXT_COMMON_WORDS
)


def parse_text_extraction_warning(value: object) -> TextExtractionWarning | None:
    if isinstance(value, str) and value in TEXT_EXTRACTION_WARNINGS:
        return value
    return None


class TextSanitizer:
    @staticmethod
    def sanitize(text: str) -> str:
        text = text.replace("\x00", "")
        return text


class TextExtractor:
    @staticmethod
    def _looks_binary(filepath: Path) -> bool:
        # A NUL byte in the leading chunk reliably marks binary content
        # (images, video, archives); UTF-8/cp1252 text never contains it.
        try:
            with open(filepath, "rb") as f:
                return b"\x00" in f.read(8192)
        except (PermissionError, OSError):
            return False

    @staticmethod
    def extract_from_plain_text(filepath: Path, filename: str | None = None) -> str:
        display_name = filename or filepath.name
        # Try UTF-8 first, then cp1252 (Windows), then UTF-8 with replacement
        # Avoid latin-1 as it accepts any byte sequence including binary garbage
        try:
            return filepath.read_text("utf-8")
        except UnicodeDecodeError:
            pass
        except PermissionError as e:
            raise ExtractionError(f"Permission denied reading '{display_name}': {e}")
        except OSError as e:
            raise ExtractionError(f"Error reading '{display_name}': {e}")

        try:
            return filepath.read_text("cp1252")
        except UnicodeDecodeError:
            pass

        # Final fallback: UTF-8 with replacement characters for undecodable bytes
        try:
            return filepath.read_text("utf-8", errors="replace")
        except (PermissionError, OSError) as e:
            raise ExtractionError(f"Error reading '{display_name}': {e}")

    @staticmethod
    def _table_to_markdown(table: list[list[str | None]]) -> str:
        rows = [
            [(cell or "").replace("\n", " ").strip() for cell in row]
            for row in table
            if any(cell for cell in row)
        ]
        if not rows:
            return ""

        lines = ["| " + " | ".join(rows[0]) + " |"]
        lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
        for row in rows[1:]:
            # Pad short rows so the markdown stays rectangular
            padded = row + [""] * (len(rows[0]) - len(row))
            lines.append("| " + " | ".join(padded[: len(rows[0])]) + " |")
        return "\n".join(lines)

    @classmethod
    def pdf_text_quality_warnings(cls, text: str) -> tuple[TextExtractionWarning, ...]:
        if cls._looks_likely_reversed_pdf_text(text):
            return (PDF_TEXT_LIKELY_REVERSED_WARNING,)
        return ()

    @classmethod
    def _looks_likely_reversed_pdf_text(cls, text: str) -> bool:
        if cls._has_vertical_letter_run(text):
            return True
        tokens = cls._text_quality_tokens(text)
        if len(tokens) < _PDF_TEXT_MIN_TOKENS_FOR_REVERSAL_WARNING:
            return False
        reversed_hits = sum(
            1 for token in tokens if token in _PDF_TEXT_REVERSED_COMMON_WORDS
        )
        if reversed_hits < _PDF_TEXT_MIN_REVERSED_COMMON_TOKENS:
            return False
        common_hits = sum(1 for token in tokens if token in _PDF_TEXT_COMMON_WORDS)
        return reversed_hits > common_hits

    @staticmethod
    def _text_quality_tokens(text: str) -> list[str]:
        tokens: list[str] = []
        token: list[str] = []
        for char in text.casefold():
            if char.isalpha():
                token.append(char)
                continue
            if token:
                tokens.append("".join(token))
                token = []
        if token:
            tokens.append("".join(token))
        return tokens

    @staticmethod
    def _has_vertical_letter_run(text: str) -> bool:
        run = 0
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if len(line) == 1 and line.isalpha():
                run += 1
                if run >= _PDF_TEXT_MIN_VERTICAL_LETTER_RUN:
                    return True
            elif line:
                run = 0
        return False

    @classmethod
    def _extract_pdf_page(cls, page: Any) -> str:
        # pdfplumber ships no type stubs — treat the page as Any and pin the
        # types we rely on at the boundaries.
        tables: list[Any] = list(page.find_tables())
        if not tables:
            return str(page.extract_text() or "")

        # Extract running text outside the table regions, then append the
        # tables as markdown — otherwise table cells appear twice.
        table_bboxes: list[tuple[float, float, float, float]] = [
            tuple(table.bbox) for table in tables
        ]

        def outside_tables(obj: dict[str, Any]) -> bool:
            center_x = (float(obj["x0"]) + float(obj["x1"])) / 2
            center_y = (float(obj["top"]) + float(obj["bottom"])) / 2
            return not any(
                x0 <= center_x <= x1 and top <= center_y <= bottom
                for (x0, top, x1, bottom) in table_bboxes
            )

        parts: list[str] = []
        text = str(page.filter(outside_tables).extract_text() or "")
        if text.strip():
            parts.append(text)
        for table in tables:
            markdown = cls._table_to_markdown(table.extract())
            if markdown:
                parts.append(markdown)
        return "\n\n".join(parts)

    @classmethod
    def extract_from_pdf(
        cls,
        filepath: Path,
        filename: str | None = None,
        *,
        limits: PdfExtractionLimits | None = None,
    ) -> str:
        display_name = filename or filepath.name
        with cls._pdf_errors(display_name):
            if limits is None:
                return cls._extract_pdf_text(filepath, display_name, limits=None)
            with cls._pdf_process(filepath, display_name, limits) as (
                process,
                result_path,
                started,
            ):
                process.join(
                    max(0, limits.timeout_seconds - (time.monotonic() - started))
                )
                return cls._pdf_process_result(process, result_path, started, limits)

    @classmethod
    async def extract_from_pdf_async(
        cls,
        filepath: Path,
        filename: str | None = None,
        *,
        limits: PdfExtractionLimits,
    ) -> str:
        display_name = filename or filepath.name
        with cls._pdf_errors(display_name):
            with cls._pdf_process(filepath, display_name, limits) as (
                process,
                result_path,
                started,
            ):
                while process.is_alive():
                    remaining = limits.timeout_seconds - (time.monotonic() - started)
                    if remaining <= 0:
                        break
                    await asyncio.sleep(min(0.05, remaining))
                return cls._pdf_process_result(process, result_path, started, limits)

    @classmethod
    @contextmanager
    def _pdf_process(
        cls, filepath: Path, display_name: str, limits: PdfExtractionLimits
    ) -> Generator[tuple[BaseProcess, Path, float]]:
        # A file avoids a full result pipe blocking child exit before join().
        with TemporaryDirectory(prefix="eneo-pdf-") as directory:
            result_path = Path(directory) / "result.json"
            started = time.monotonic()
            process = multiprocessing.get_context("spawn").Process(
                target=cls._extract_pdf_in_child,
                args=(
                    filepath,
                    display_name,
                    limits,
                    started + limits.timeout_seconds,
                    result_path,
                ),
            )
            try:
                process.start()
                yield process, result_path, started
            finally:
                if process.is_alive():
                    process.kill()
                if process.pid is not None:
                    process.join()
                process.close()

    @staticmethod
    def _pdf_process_result(
        process: BaseProcess,
        result_path: Path,
        started: float,
        limits: PdfExtractionLimits,
    ) -> str:
        elapsed = time.monotonic() - started
        if (
            process.is_alive()
            or elapsed >= limits.timeout_seconds
            or process.exitcode == -signal.SIGALRM
        ):
            raise PdfExtractionLimitExceeded(
                "seconds", max(elapsed, limits.timeout_seconds), limits.timeout_seconds
            )
        if process.exitcode != 0:
            raise ExtractionError(
                f"PDF extraction process exited with code {process.exitcode}"
            )
        result = _PDF_RESULT.validate_json(result_path.read_bytes())
        if isinstance(result, str):
            return result
        if isinstance(result, _PdfLimitResult):
            raise PdfExtractionLimitExceeded(
                result.limit, result.measured, result.ceiling
            )
        if result.kind == "encrypted":
            raise PDFPasswordIncorrect(result.details)
        if result.kind == "corrupt":
            raise PDFSyntaxError(result.details)
        raise ExtractionError(result.details)

    @staticmethod
    @contextmanager
    def _pdf_errors(display_name: str) -> Generator[None]:
        try:
            yield
        except ExtractionError:
            raise
        except PDFPasswordIncorrect as e:
            logger.warning(f"Password-protected PDF rejected: {display_name}")
            raise EncryptedFileError(display_name) from e
        except PDFSyntaxError as e:
            logger.warning(f"PDF read error for {display_name}: {e}")
            raise CorruptFileError(display_name, str(e))
        except Exception as e:
            logger.error(f"Unexpected PDF extraction error for {display_name}: {e}")
            raise ExtractionError(
                f"PDF extraction failed for '{display_name}': {str(e)}"
            )

    @classmethod
    def _extract_pdf_in_child(
        cls,
        filepath: Path,
        display_name: str,
        limits: PdfExtractionLimits,
        deadline: float,
        result_path: Path,
    ) -> None:
        # The OS deadline survives supervisor death and does not require the
        # parser to release the GIL or execute a Python signal handler.
        signal.signal(signal.SIGALRM, signal.SIG_DFL)
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            signal.raise_signal(signal.SIGALRM)
        signal.setitimer(signal.ITIMER_REAL, remaining)
        result: str | _PdfLimitResult | _PdfErrorResult
        try:
            result = cls._extract_pdf_text(filepath, display_name, limits)
        except PdfExtractionLimitExceeded as exc:
            result = _PdfLimitResult(
                limit=exc.limit, measured=exc.measured, ceiling=exc.ceiling
            )
        except PDFPasswordIncorrect as exc:
            result = _PdfErrorResult(kind="encrypted", details=str(exc))
        except PDFSyntaxError as exc:
            result = _PdfErrorResult(kind="corrupt", details=str(exc))
        except Exception as exc:
            result = _PdfErrorResult(
                kind="error",
                details=f"PDF extraction failed for '{display_name}': {exc}",
            )
        result_path.write_bytes(_PDF_RESULT.dump_json(result))

    @classmethod
    def _extract_pdf_text(
        cls, filepath: Path, display_name: str, limits: PdfExtractionLimits | None
    ) -> str:
        with pdfplumber.open(filepath) as pdf:
            if limits is not None and len(pdf.pages) > limits.max_pages:
                raise PdfExtractionLimitExceeded(
                    "pages", len(pdf.pages), limits.max_pages
                )
            page_texts: list[str] = []
            extracted_bytes = 0
            has_content = False
            for page in pdf.pages:
                try:
                    page_text = cls._extract_pdf_page(page)
                except PdfExtractionLimitExceeded:
                    raise
                except Exception as e:
                    logger.warning(
                        f"Table-aware extraction failed on page "
                        f"{page.page_number} of '{display_name}', "
                        f"falling back to plain text: {e}"
                    )
                    page_text = page.extract_text() or ""
                framed_text = f"[PAGE {page.page_number}]\n{page_text}"
                if limits is not None:
                    extracted_bytes += len(framed_text.encode("utf-8"))
                    if page_texts:
                        extracted_bytes += 2
                    if extracted_bytes > limits.max_extracted_bytes:
                        raise PdfExtractionLimitExceeded(
                            "extracted_bytes",
                            extracted_bytes,
                            limits.max_extracted_bytes,
                        )
                has_content = has_content or bool(page_text.strip())
                page_texts.append(framed_text)

            extracted_text = "\n\n".join(page_texts)

        if not has_content:
            logger.warning(
                f"No text extracted from PDF '{display_name}' - "
                "file may be image-only or scanned"
            )
            # Bare page markers must not bypass downstream empty-content handling.
            return ""

        return TextSanitizer.sanitize(extracted_text)

    @staticmethod
    def extract_from_docx(filepath: Path, filename: str | None = None) -> str:
        display_name = filename or filepath.name
        try:
            with docx2python(filepath) as docx_content:
                return docx_content.text
        except zipfile.BadZipFile:
            raise CorruptFileError(
                display_name,
                "Invalid ZIP structure - file may be corrupted or in legacy .doc format",
            )
        except KeyError as e:
            raise CorruptFileError(
                display_name, f"Missing required document component: {e}"
            )
        except Exception as e:
            logger.error(f"Unexpected DOCX extraction error for {display_name}: {e}")
            raise ExtractionError(
                f"DOCX extraction failed for '{display_name}': {str(e)}"
            )

    @staticmethod
    def extract_from_xlsx(filepath: Path, filename: str | None = None) -> str:
        import pandas as pd

        display_name = filename or filepath.name
        try:
            # pandas-stubs miss the context-manager protocol on ExcelFile, so
            # close() explicitly to avoid leaking the underlying file handle.
            xls = pd.ExcelFile(filepath, engine="calamine")
            try:
                parts: list[str] = []

                # Global file context (helpful for first chunk and direct chat)
                parts.append(f"File: {display_name}")

                for sheet_name in xls.sheet_names:
                    df: pd.DataFrame = pd.read_excel(  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType]  # pandas stubs are incomplete
                        xls, sheet_name=sheet_name, engine="calamine"
                    )

                    if df.empty:  # pyright: ignore[reportUnknownMemberType]  # pandas stubs are incomplete
                        continue

                    # Handle merged cells - forward fill values
                    df = df.ffill()  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType]  # pandas stubs are incomplete

                    # Clean column names (ensure strings, remove newlines)
                    df.columns = (  # pyright: ignore[reportUnknownMemberType]  # pandas stubs are incomplete
                        df.columns.astype(str).str.replace("\n", " ")  # pyright: ignore[reportUnknownMemberType]  # pandas stubs are incomplete
                    )

                    # Serialize each row as self-contained key-value pairs
                    # This ensures every chunk has full context, even if split
                    def serialize_row(row: pd.Series) -> str:  # type: ignore[type-arg]  # pd.Series generic param unavailable at runtime
                        # Filter out NaN to save tokens
                        pairs = [
                            f"{col}: {val}"
                            for col, val in row.items()
                            if pd.notna(val)  # pyright: ignore[reportUnknownArgumentType, reportUnknownMemberType]  # pandas stubs are incomplete
                        ]
                        return f"Sheet: {sheet_name} | " + " | ".join(pairs)

                    # pandas apply/str.cat return types are unknown due to incomplete stubs
                    raw_sheet_text = df.apply(serialize_row, axis=1).str.cat(sep="\n")  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType]  # pandas stubs are incomplete
                    sheet_text: str = str(raw_sheet_text)  # pyright: ignore[reportUnknownArgumentType]  # pandas stubs are incomplete
                    parts.append(sheet_text)

                return "\n\n".join(parts)
            finally:
                xls.close()  # pyright: ignore[reportUnknownMemberType, reportAttributeAccessIssue]  # pandas-stubs miss ExcelFile.close
        except ValueError as e:
            raise CorruptFileError(display_name, f"Cannot parse Excel format: {e}")
        except Exception as e:
            logger.error(f"Unexpected Excel extraction error for {display_name}: {e}")
            raise ExtractionError(
                f"Excel extraction failed for '{display_name}': {str(e)}"
            )

    @staticmethod
    def extract_from_pptx(filepath: Path, filename: str | None = None) -> str:
        display_name = filename or filepath.name
        try:
            # Extract text from pptx using python-pptx
            # Use list join instead of string concatenation for O(n) vs O(n^2) complexity
            # pptx.Presentation accepts str, not Path — convert explicitly
            presentation = pptx.Presentation(str(filepath))
            parts: list[str] = []
            for slide in presentation.slides:
                slide_parts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        # Collect all text from runs in this shape
                        # python-pptx stubs are incomplete; member types are partially unknown
                        shape_text = " ".join(
                            run.text  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType]  # pptx stubs are incomplete
                            for para in shape.text_frame.paragraphs  # type: ignore[attr-defined]
                            for run in para.runs  # pyright: ignore[reportUnknownVariableType, reportUnknownMemberType]  # pptx stubs are incomplete
                            if run.text  # pyright: ignore[reportUnknownMemberType]  # pptx stubs are incomplete
                        )
                        if shape_text.strip():
                            slide_parts.append(shape_text)
                if slide.has_notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType, reportOptionalMemberAccess]  # pptx stubs are incomplete
                    notes_text = (
                        (notes_frame.text or "").strip()
                        if notes_frame is not None
                        else ""
                    )  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType]  # pptx stubs are incomplete
                    if notes_text:
                        slide_parts.append(f"Speaker notes: {notes_text}")
                if slide_parts:
                    parts.append(" ".join(slide_parts))
            return "\n".join(parts)
        except (zipfile.BadZipFile, PackageNotFoundError):
            raise CorruptFileError(
                display_name,
                "Invalid ZIP structure - file may be corrupted or in legacy .ppt format",
            )
        except KeyError as e:
            raise CorruptFileError(
                display_name, f"Missing required presentation component: {e}"
            )
        except Exception as e:
            logger.error(f"Unexpected PPTX extraction error for {display_name}: {e}")
            raise ExtractionError(
                f"PPTX extraction failed for '{display_name}': {str(e)}"
            )

    def extract(
        self,
        filepath: Path,
        mimetype: str | None = None,
        filename: str | None = None,
        *,
        pdf_limits: PdfExtractionLimits | None = None,
    ) -> str:
        mimetype = mimetype or magic.from_file(filepath, mime=True)  # pyright: ignore[reportUnknownMemberType]  # python-magic stubs are incomplete
        # Use original filename for error messages, fallback to temp filepath
        display_name = filename or filepath.name

        # Reject legacy formats early with helpful message
        if mimetype == TextMimeTypes.DOC.value:
            raise UnsupportedFormatError(
                display_name,
                ".doc (Legacy Word) - please save as .docx",
            )
        if mimetype == TextMimeTypes.PPT.value:
            raise UnsupportedFormatError(
                display_name,
                ".ppt (Legacy PowerPoint) - please save as .pptx",
            )

        match mimetype:
            case (
                TextMimeTypes.TXT
                | TextMimeTypes.MD
                | TextMimeTypes.TEXT_CSV
                | TextMimeTypes.APP_CSV
                | TextMimeTypes.JSON
                | TextMimeTypes.XML
                | TextMimeTypes.XML_APP
            ):
                extracted_text = self.extract_from_plain_text(filepath, display_name)
            case TextMimeTypes.PDF:
                extracted_text = self.extract_from_pdf(
                    filepath, display_name, limits=pdf_limits
                )
            case TextMimeTypes.DOCX:
                extracted_text = self.extract_from_docx(filepath, display_name)
            case TextMimeTypes.PPTX:
                extracted_text = self.extract_from_pptx(filepath, display_name)
            case TextMimeTypes.XLSX | TextMimeTypes.XLS:
                extracted_text = self.extract_from_xlsx(filepath, display_name)
            case _:
                # Unknown mimetype: only treat as text if it is not binary.
                # Guards against binary uploads (e.g. unsupported image formats)
                # being decoded into garbage and written to the TEXT column.
                if self._looks_binary(filepath):
                    raise UnsupportedFormatError(display_name, mimetype or "binary")
                extracted_text = self.extract_from_plain_text(filepath, display_name)

        return extracted_text.strip()
