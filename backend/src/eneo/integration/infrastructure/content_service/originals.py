"""Which synced originals are kept, and the schema stub for tool-only tables.

A synced document normally survives only as extracted text. Tabular files are
the exception: a spreadsheet's value lies in computing over it, which needs
the bytes, so their originals are retained and handed to analysis tools as
signed references. This module is the one place that decides which files
that applies to.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator, Mapping

from openpyxl import load_workbook

# Extension -> media type of the originals a sync keeps. Growing this list is
# the whole change needed to hand another kind of file to tools.
RETAINED_ORIGINALS: Mapping[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "tsv": "text/tab-separated-values",
}
_EXTENSION_BY_MEDIA_TYPE = {media: ext for ext, media in RETAINED_ORIGINALS.items()}
TABULAR_EXTENSIONS = frozenset({"xlsx", "csv", "tsv"})
PLACEHOLDER_MAX_CHARS = 2000
_PLACEHOLDER_MAX_SHEETS = 20


def _extension(filename: str) -> str:
    name = filename.rsplit("/", 1)[-1]
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def original_media_type(filename: str, content_type: str | None = None) -> str | None:
    """The media type to store the original under, or None when it is not kept.

    The file name decides; the response header only helps when the name has
    no extension.
    """
    extension = _extension(filename)
    if extension:
        return RETAINED_ORIGINALS.get(extension)
    if content_type:
        base = content_type.split(";", 1)[0].strip().lower()
        if base in _EXTENSION_BY_MEDIA_TYPE:
            return base
    return None


def retain_original_bytes(filename: str, content_type: str | None = None) -> bool:
    return original_media_type(filename, content_type) is not None


def is_tabular(filename: str, content_type: str | None = None) -> bool:
    """Whether the source's tabular policy applies to this file."""
    media_type = original_media_type(filename, content_type)
    return media_type is not None and _EXTENSION_BY_MEDIA_TYPE[media_type] in (
        TABULAR_EXTENSIONS
    )


def build_tabular_placeholder_text(filename: str, data: bytes) -> str:
    """Title, sheet names and header rows of a table kept for tools only.

    Enough for a reader to know what the file holds and that the table tools
    open it; never the rows themselves, which are not indexed.
    """
    lines = [
        f"Spreadsheet: {filename}",
        "Rows are not indexed for search; the file is available to the table "
        "analysis tools.",
        "",
    ]
    try:
        for sheet, columns in _sheet_headers(filename, data):
            header = ", ".join(str(column) for column in columns) or "(no header row)"
            lines.append(
                f'Sheet "{sheet}": {header}' if sheet else f"Columns: {header}"
            )
    except Exception:  # noqa: BLE001 - a header preview must never fail a sync
        lines.append("(column headers could not be read)")
    return "\n".join(lines)[:PLACEHOLDER_MAX_CHARS]


def _sheet_headers(
    filename: str, data: bytes
) -> Iterator[tuple[str | None, list[str]]]:
    if _extension(filename) == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets[:_PLACEHOLDER_MAX_SHEETS]:
                first_row = next(
                    sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
                )
                yield sheet.title, [str(cell) for cell in first_row if cell is not None]
        finally:
            workbook.close()
        return
    delimiter = "\t" if _extension(filename) == "tsv" else ","
    first_line = (
        data.split(b"\n", 1)[0].decode("utf-8", errors="replace").lstrip("\ufeff")
    )
    no_header: list[str] = []
    yield None, next(csv.reader([first_line], delimiter=delimiter), no_header)
