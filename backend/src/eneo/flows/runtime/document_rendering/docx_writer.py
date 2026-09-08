"""Write ``DocumentBlock`` content into a DOCX with the target document's own styles.

The writer is the single owner of "blocks -> WordprocessingML". It resolves
every style by name against the document it writes into (Heading N, the
list styles, the table style), marks table header rows, restarts numbering
per list and never applies direct paragraph formatting, so a template's
look and its accessibility properties survive whatever the model wrote.
Both DOCX product modes use it: verbatim rendering fills the standard body
template, template fill writes into each content control.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Sequence, cast

from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.run import CT_R
from lxml.etree import Element

from eneo.flows.runtime.document_rendering.blocks import (
    MAX_HEADING_LEVEL,
    DocumentBlock,
    DocumentStructureError,
    InlineRuns,
    InlineTextRun,
)

_MONOSPACE_FONT = "Consolas"
_LIST_BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")
_LIST_NUMBER_STYLES = ("List Number", "List Number 2", "List Number 3")


def word_element(tag: str) -> Element:
    """Create registered Word XML with the common lxml element interface.

    A document contains both registered python-docx subclasses and generic
    lxml elements, including content controls. Parents must accept both.
    With postponed annotations, Element uses types-lxml's public element alias;
    OxmlElement remains the runtime factory that registers Word-specific types.
    """
    return OxmlElement(tag)


@dataclass(frozen=True, slots=True)
class _ListStyles:
    bullet: tuple[str, ...]
    number: tuple[str, ...]


class DocxBlockWriter:
    def __init__(self, document: Any) -> None:
        self._document = document
        self._style_names = {style.name for style in document.styles}
        self._table_style = self._resolve_table_style()
        self._list_styles = _ListStyles(
            bullet=self._resolve_ladder(_LIST_BULLET_STYLES),
            number=self._resolve_ladder(_LIST_NUMBER_STYLES),
        )
        self._code_style = self._first_existing("Code", "No Spacing")
        self._hyperlink_style = self._first_existing("Hyperlink")

    def elements(
        self,
        blocks: Sequence[DocumentBlock],
        *,
        heading_base: int = 0,
        section_label: str = "the document",
    ) -> list[Any]:
        """Build the block elements for ``blocks`` in document order.

        ``heading_base`` is the outline level the content sits under (0 for a
        whole document, 1 for a section under a Heading 1). The shallowest
        heading in the content becomes ``heading_base + 1`` and deeper
        headings follow relatively, so a model that writes ``##`` and one that
        writes ``#`` for its sub-sections produce the same outline.
        """

        levels = _relative_heading_levels(
            blocks, heading_base=heading_base, section_label=section_label
        )
        elements: list[Any] = []
        heading_index = 0
        for block in blocks:
            if block.kind == "empty":
                continue
            if block.kind == "heading":
                level = levels[heading_index]
                heading_index += 1
                elements.append(self._heading(block, level=level))
            elif block.kind == "bullet_list":
                elements.extend(self._list(block, ordered=False))
            elif block.kind == "numbered_list":
                elements.extend(self._list(block, ordered=True))
            elif block.kind == "code":
                elements.append(self._code(block))
            elif block.kind == "table":
                table = self._table(block)
                if table is not None:
                    elements.append(table)
            else:
                elements.append(self._paragraph(block))
        return elements

    # -- block builders -------------------------------------------------------

    def _heading(self, block: DocumentBlock, *, level: int) -> Any:
        style_name = f"Heading {level}"
        if style_name not in self._style_names:
            raise DocumentStructureError(
                f"The document template has no '{style_name}' style; use at most "
                f"{self._deepest_heading_level()} heading levels."
            )
        paragraph = self._document.add_paragraph(style=style_name)
        self._append_runs(paragraph, block.runs, fallback=block.text)
        return paragraph._p

    def _paragraph(self, block: DocumentBlock) -> Any:
        paragraph = self._document.add_paragraph()
        self._append_runs(paragraph, block.runs, fallback=block.text)
        return paragraph._p

    def _code(self, block: DocumentBlock) -> Any:
        paragraph = (
            self._document.add_paragraph(style=self._code_style)
            if self._code_style is not None
            else self._document.add_paragraph()
        )
        run = paragraph.add_run(block.text)
        run.font.name = _MONOSPACE_FONT
        return paragraph._p

    def _list(self, block: DocumentBlock, *, ordered: bool) -> list[Any]:
        ladder = self._list_styles.number if ordered else self._list_styles.bullet
        numbering_by_level: list[int | None] = []
        elements: list[Any] = []
        for index, item in enumerate(block.items):
            level = block.item_level(index)
            if level >= len(ladder):
                raise DocumentStructureError(
                    f"The document template has no list style for level {level + 1}."
                )
            paragraph = self._document.add_paragraph(style=ladder[level])
            if ordered:
                del numbering_by_level[level + 1 :]
                while len(numbering_by_level) <= level:
                    depth = len(numbering_by_level)
                    numbering_by_level.append(
                        self._restart_numbering(
                            ladder[depth], start=block.start if depth == 0 else 1
                        )
                    )
                num_id = numbering_by_level[level]
                if num_id is not None:
                    _set_numbering(paragraph, num_id=num_id, level=0)
            runs = block.item_runs[index] if index < len(block.item_runs) else ()
            self._append_runs(paragraph, runs, fallback=item)
            elements.append(paragraph._p)
        return elements

    def _table(self, block: DocumentBlock) -> Any | None:
        if not block.rows:
            return None
        column_count = max(len(row) for row in block.rows)
        table = self._document.add_table(rows=len(block.rows), cols=column_count)
        if self._table_style is not None:
            table.style = self._document.styles[self._table_style]
        _mark_header_row(table)
        header_needs_emphasis = self._table_style in (None, "Table Grid")
        for row_index, row in enumerate(table.rows):
            values = block.rows[row_index]
            value_runs = (
                block.row_runs[row_index] if row_index < len(block.row_runs) else ()
            )
            for column_index, cell in enumerate(row.cells):
                paragraph = cell.paragraphs[0]
                runs = (
                    value_runs[column_index] if column_index < len(value_runs) else ()
                )
                fallback = values[column_index] if column_index < len(values) else ""
                self._append_runs(
                    paragraph,
                    runs,
                    fallback=fallback,
                    force_bold=row_index == 0 and header_needs_emphasis,
                )
        return table._tbl

    # -- inline content -------------------------------------------------------

    def _append_runs(
        self,
        paragraph: Any,
        runs: InlineRuns,
        *,
        fallback: str,
        force_bold: bool = False,
    ) -> None:
        if not runs:
            if fallback:
                run = paragraph.add_run(fallback)
                if force_bold:
                    run.bold = True
            return
        for inline_run in runs:
            self._append_run(paragraph, inline_run, force_bold=force_bold)

    def _append_run(
        self, paragraph: Any, inline_run: InlineTextRun, *, force_bold: bool
    ) -> None:
        if inline_run.href:
            run = self._hyperlink_run(paragraph, inline_run)
        else:
            run = paragraph.add_run(inline_run.text)
        # Absent emphasis stays inherited from the style; only explicit
        # emphasis is written, so a bold heading style keeps its weight.
        if inline_run.bold or force_bold:
            run.bold = True
        if inline_run.italic:
            run.italic = True
        if inline_run.strikethrough:
            run.font.strike = True
        if inline_run.code:
            run.font.name = _MONOSPACE_FONT

    def _hyperlink_run(self, paragraph: Any, inline_run: InlineTextRun) -> Any:
        from docx.text.run import Run

        relationship_id = paragraph.part.relate_to(
            inline_run.href, RT.HYPERLINK, is_external=True
        )
        hyperlink = word_element("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run_element = word_element("w:r")
        text = word_element("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = inline_run.text
        run_element.append(text)
        hyperlink.append(run_element)
        paragraph._p.append(hyperlink)
        run = Run(cast(CT_R, run_element), paragraph)
        if self._hyperlink_style is not None:
            run.style = self._document.styles[self._hyperlink_style]
        return run

    # -- style resolution -----------------------------------------------------

    def _first_existing(self, *candidates: str) -> str | None:
        for candidate in candidates:
            if candidate in self._style_names:
                return candidate
        return None

    def _resolve_ladder(self, candidates: tuple[str, ...]) -> tuple[str, ...]:
        """Keep the leading run of existing list styles."""

        ladder: list[str] = []
        for candidate in candidates:
            if candidate not in self._style_names:
                break
            ladder.append(candidate)
        return tuple(ladder)

    def _resolve_table_style(self) -> str | None:
        default = self._document.styles.default(WD_STYLE_TYPE.TABLE)
        if default is not None and default.name != "Normal Table":
            return str(default.name)
        return self._first_existing("Table Grid")

    def _deepest_heading_level(self) -> int:
        deepest = 0
        for level in range(1, MAX_HEADING_LEVEL + 1):
            if f"Heading {level}" in self._style_names:
                deepest = level
        return deepest

    def _restart_numbering(self, style_name: str, *, start: int) -> int | None:
        """Give this list its own numbering instance so it starts at ``start``.

        Word continues a numbered list across paragraphs that share the style's
        numbering instance, so two separate lists would read 1, 2, 3, 4, 5.
        A fresh ``w:num`` with a start override restarts the count.
        """

        style = self._document.styles[style_name]
        style_num_pr = (
            style.element.pPr.numPr if style.element.pPr is not None else None
        )
        if style_num_pr is None or style_num_pr.numId is None:
            return None
        numbering = self._document.part.numbering_part.element
        source_num = numbering.num_having_numId(style_num_pr.numId.val)
        if source_num is None:
            return None
        new_num = numbering.add_num(source_num.abstractNumId.val)
        new_num.add_lvlOverride(ilvl=0).add_startOverride(start)
        return int(new_num.numId)


def _relative_heading_levels(
    blocks: Sequence[DocumentBlock],
    *,
    heading_base: int,
    section_label: str,
) -> list[int]:
    """Map source heading levels to document levels and validate the outline."""

    source_levels = [max(block.level, 1) for block in blocks if block.kind == "heading"]
    if not source_levels:
        return []
    shallowest = min(source_levels)
    levels: list[int] = []
    previous = heading_base
    for source_level in source_levels:
        level = heading_base + 1 + (source_level - shallowest)
        if level > previous + 1:
            raise DocumentStructureError(
                f"The headings in {section_label} skip a level (a level "
                f"{level - heading_base} heading follows level "
                f"{previous - heading_base}); use each heading level in order."
            )
        if level > MAX_HEADING_LEVEL:
            raise DocumentStructureError(
                f"The headings in {section_label} go deeper than the "
                f"{MAX_HEADING_LEVEL} levels a document supports."
            )
        levels.append(level)
        previous = level
    return levels


def _mark_header_row(table: Any) -> None:
    """Flag row 0 as the header row (repeats across pages, read as headers)."""

    if not table.rows:
        return
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    if row_properties.find(qn("w:tblHeader")) is None:
        row_properties.append(word_element("w:tblHeader"))
    table_properties = table._tbl.tblPr
    look = table_properties.find(qn("w:tblLook"))
    if look is None:
        look = word_element("w:tblLook")
        table_properties.append(look)
    look.set(qn("w:firstRow"), "1")


def _set_numbering(paragraph: Any, *, num_id: int, level: int) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = num_id
    num_pr.get_or_add_ilvl().val = level
