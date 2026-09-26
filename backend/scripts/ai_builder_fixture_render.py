"""Battle fixture content specs and one renderer per format.

A content spec (``fixtures/ai_builder_battle/specs/<file>.json``) describes one document the corpus hands to the
Builder: the real document type it imitates, then its content. ``parse_spec`` validates it (closed keys) and
``render`` turns it into bytes. Every spec has ``file`` (its own name minus ``.json``), ``format``, ``source``
({catalogue_id, url}) and a one-line ``description``, plus by format:

- pdf, docx, txt: ``letterhead`` {organisation, unit, address?}, optional ``date`` (YYYY-MM-DD) and
  ``diarienummer``, ``title`` and ordered ``blocks``: {heading}, {paragraph}, {list: [...]},
  {table: {columns, rows}} (text cells), {signature: {name, role}}, {page_break: true} (pdf and docx only), and
  in a docx template {control: {kind, tag, label, hint}}: a Word content control as the product's template runtime
  reads it (kind rich, a body-level section, or text, one line in its own paragraph; tags unique);
- csv (exactly one sheet), xlsx: ``sheets`` [{name, columns, rows}], names as Excel allows them and distinct
  ignoring case, cells strings, finite numbers or null;
- json: ``data`` (finite numbers only), written in the spec's key order.

Rendering depends on nothing outside the spec but the installed libraries and the fonts WeasyPrint finds, and every
renderer pins its timestamps, so one machine renders a spec to the same bytes every time. Another machine's fonts can
change the PDF bytes and page layout, which is why the drift check compares extracted text, not bytes.
"""

from __future__ import annotations

import csv
import html
import io
import json
import math
import zipfile
from collections.abc import Callable, Collection, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Literal, TypeVar, cast

import docx
import weasyprint  # pyright: ignore[reportMissingTypeStubs]  # WeasyPrint publishes no stubs
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.styles.style import ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from eneo.flows.runtime.document_rendering.docx_content_controls import (
    append_rich_control,
    append_text_control,
)

Format = Literal["pdf", "docx", "xlsx", "csv", "json", "txt"]
FORMATS: tuple[Format, ...] = ("pdf", "docx", "xlsx", "csv", "json", "txt")
FIXED_TIMESTAMP = datetime(2026, 1, 1, 0, 0, 0)
FIXED_ZIP_TIMESTAMP = (2026, 1, 1, 0, 0, 0)
AUTHOR = "Eneo battle corpus"

Cell = str | int | float | None
_T = TypeVar("_T")


@dataclass(frozen=True, slots=True)
class Source:
    catalogue_id: str
    url: str


@dataclass(frozen=True, slots=True)
class Letterhead:
    organisation: str
    unit: str
    address: str | None = None


@dataclass(frozen=True, slots=True)
class Heading:
    text: str


@dataclass(frozen=True, slots=True)
class Paragraph:
    text: str


@dataclass(frozen=True, slots=True)
class BulletList:
    items: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class Table:
    columns: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True, slots=True)
class Signature:
    name: str
    role: str


@dataclass(frozen=True, slots=True)
class PageBreak:
    pass


@dataclass(frozen=True, slots=True)
class Control:
    """A Word content control, the fill target of a DOCX template: rich spans paragraphs, text is one line."""

    kind: Literal["rich", "text"]
    tag: str
    label: str
    hint: str


Block = Heading | Paragraph | BulletList | Table | Signature | PageBreak | Control


@dataclass(frozen=True, slots=True)
class DocumentContent:
    letterhead: Letterhead
    date: str | None
    diarienummer: str | None
    title: str
    blocks: tuple[Block, ...]

    def metadata(self) -> list[tuple[str, str]]:
        """The dated, registered identity every document renders above its title."""

        return [
            (label, value)
            for label, value in (
                ("Datum", self.date),
                ("Diarienummer", self.diarienummer),
            )
            if value is not None
        ]


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    columns: tuple[str, ...]
    rows: tuple[tuple[Cell, ...], ...]


@dataclass(frozen=True, slots=True)
class JsonData:
    value: object


@dataclass(frozen=True, slots=True)
class FixtureSpec:
    file: str
    format: Format
    source: Source
    description: str
    content: DocumentContent | tuple[Sheet, ...] | JsonData


_COMMON_KEYS = ("file", "format", "source", "description")
_DOCUMENT_KEYS = ("letterhead", "title", "blocks")
_CONTENT_KEYS: dict[Format, tuple[str, ...]] = {
    "pdf": _DOCUMENT_KEYS,
    "docx": _DOCUMENT_KEYS,
    "txt": _DOCUMENT_KEYS,
    "xlsx": ("sheets",),
    "csv": ("sheets",),
    "json": ("data",),
}
_BLOCK_KINDS = (
    "heading",
    "paragraph",
    "list",
    "table",
    "signature",
    "page_break",
    "control",
)


def _fields(
    value: object,
    where: str,
    required: Collection[str],
    optional: Collection[str] = (),
) -> Mapping[str, object]:
    if not isinstance(value, Mapping):
        raise ValueError(f"{where} must be an object")
    fields = cast(Mapping[str, object], value)
    missing = sorted(set(required) - set(fields))
    if missing:
        raise ValueError(f"{where}: missing keys: {', '.join(missing)}")
    unknown = sorted(set(fields) - set(required) - set(optional))
    if unknown:
        raise ValueError(f"{where}: unknown keys: {', '.join(unknown)}")
    return fields


def _text(value: object, where: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{where} must be a non-empty string")
    return value


def _optional_text(value: object, where: str) -> str | None:
    return None if value is None else _text(value, where)


def _items(value: object, where: str) -> list[object]:
    if not isinstance(value, list) or not value:
        raise ValueError(f"{where} must be a non-empty list")
    return cast(list[object], value)


def _rows(
    value: object,
    where: str,
    width: int,
    cell: Callable[[object, str], _T],
) -> tuple[tuple[_T, ...], ...]:
    rows: list[tuple[_T, ...]] = []
    for number, raw_row in enumerate(_items(value, f"{where} rows"), start=1):
        row = _items(raw_row, f"{where} row {number}")
        if len(row) != width:
            raise ValueError(
                f"{where} row {number} has {len(row)} cells for {width} columns"
            )
        rows.append(tuple(cell(item, f"{where} row {number}") for item in row))
    return tuple(rows)


def _table_cell(value: object, where: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{where}: table cells must be strings")
    return value


def _sheet_cell(value: object, where: str) -> Cell:
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{where}: a cell must be a finite number")
    if value is None or (
        isinstance(value, str | int | float) and not isinstance(value, bool)
    ):
        return value
    raise ValueError(f"{where}: a cell must be a string, a number or null")


def _finite(value: object) -> bool:
    if isinstance(value, float):
        return math.isfinite(value)
    if isinstance(value, list):
        return all(_finite(item) for item in cast(list[object], value))
    if isinstance(value, Mapping):
        return all(_finite(item) for item in cast(Mapping[str, object], value).values())
    return True


def _block(value: object, where: str, fmt: Format) -> Block:
    if not isinstance(value, Mapping) or len(cast(Mapping[str, object], value)) != 1:
        raise ValueError(f"{where} must have exactly one of {', '.join(_BLOCK_KINDS)}")
    ((kind, body),) = cast(Mapping[str, object], value).items()
    match kind:
        case "heading":
            return Heading(_text(body, f"{where}.heading"))
        case "paragraph":
            return Paragraph(_text(body, f"{where}.paragraph"))
        case "list":
            items = _items(body, f"{where}.list")
            return BulletList(tuple(_text(item, f"{where}.list") for item in items))
        case "table":
            fields = _fields(body, f"{where}.table", ("columns", "rows"))
            columns = tuple(
                _table_cell(column, f"{where}.table columns")
                for column in _items(fields["columns"], f"{where}.table columns")
            )
            rows = _rows(fields["rows"], f"{where}.table", len(columns), _table_cell)
            return Table(columns, rows)
        case "signature":
            fields = _fields(body, f"{where}.signature", ("name", "role"))
            return Signature(
                _text(fields["name"], f"{where}.signature name"),
                _text(fields["role"], f"{where}.signature role"),
            )
        case "page_break":
            if body is not True:
                raise ValueError(f"{where}: page_break must be true")
            if fmt == "txt":
                raise ValueError(f"{where}: a txt document has no page_break")
            return PageBreak()
        case "control":
            if fmt != "docx":
                raise ValueError(f"{where}: only a docx has content controls")
            fields = _fields(body, f"{where}.control", ("kind", "tag", "label", "hint"))
            kind = fields["kind"]
            if kind not in ("rich", "text"):
                raise ValueError(f"{where}: control kind must be rich or text")
            return Control(
                kind,
                _text(fields["tag"], f"{where}.control tag"),
                _text(fields["label"], f"{where}.control label"),
                _text(fields["hint"], f"{where}.control hint"),
            )
        case _:
            raise ValueError(
                f"{where} must have exactly one of {', '.join(_BLOCK_KINDS)}"
            )


def _date(value: object, where: str) -> str | None:
    text = _optional_text(value, where)
    if text is None:
        return None
    try:
        if date.fromisoformat(text).isoformat() == text:
            return text
    except ValueError:
        pass
    raise ValueError(f"{where} must be YYYY-MM-DD")


def _document(fields: Mapping[str, object], where: str, fmt: Format) -> DocumentContent:
    letterhead = _fields(
        fields["letterhead"],
        f"{where}: letterhead",
        ("organisation", "unit"),
        ("address",),
    )
    blocks = tuple(
        _block(block, f"{where}: blocks[{index}]", fmt)
        for index, block in enumerate(_items(fields["blocks"], f"{where}: blocks"))
    )
    tags = [block.tag for block in blocks if isinstance(block, Control)]
    if len(set(tags)) != len(tags):
        raise ValueError(f"{where}: control tags must be unique")
    return DocumentContent(
        letterhead=Letterhead(
            _text(letterhead["organisation"], f"{where}: letterhead organisation"),
            _text(letterhead["unit"], f"{where}: letterhead unit"),
            _optional_text(letterhead.get("address"), f"{where}: letterhead address"),
        ),
        date=_date(fields.get("date"), f"{where}: date"),
        diarienummer=_optional_text(
            fields.get("diarienummer"), f"{where}: diarienummer"
        ),
        title=_text(fields["title"], f"{where}: title"),
        blocks=blocks,
    )


def _sheets(value: object, where: str, fmt: Format) -> tuple[Sheet, ...]:
    raw_sheets = _items(value, f"{where}: sheets")
    if fmt == "csv" and len(raw_sheets) != 1:
        raise ValueError(f"{where}: a csv has exactly one sheet")
    sheets: list[Sheet] = []
    for index, raw_sheet in enumerate(raw_sheets):
        label = f"{where}: sheets[{index}]"
        fields = _fields(raw_sheet, label, ("name", "columns", "rows"))
        columns = tuple(
            _text(column, f"{label} columns")
            for column in _items(fields["columns"], label)
        )
        if len(set(columns)) != len(columns):
            raise ValueError(f"{label}: columns must be unique non-empty strings")
        name = _text(fields["name"], f"{label} name")
        if len(name) > 31 or set(name) & set("[]:*?/\\"):
            raise ValueError(
                f"{label}: a sheet name has at most 31 characters and none of [ ] : * ? / \\"
            )
        sheets.append(
            Sheet(
                name,
                columns,
                _rows(fields["rows"], label, len(columns), _sheet_cell),
            )
        )
    if len({sheet.name.casefold() for sheet in sheets}) != len(sheets):
        # openpyxl would silently rename the second of two names that differ only in case.
        raise ValueError(f"{where}: sheet names must differ in more than case")
    return tuple(sheets)


def parse_spec(raw: object, *, spec_name: str) -> FixtureSpec:
    """Validate one content spec; every error names the spec file."""

    if not isinstance(raw, Mapping):
        raise ValueError(f"{spec_name} must be an object")
    raw_fields = cast(Mapping[str, object], raw)
    fmt = raw_fields.get("format")
    if fmt not in FORMATS:
        raise ValueError(f"{spec_name}: format must be one of {', '.join(FORMATS)}")
    optional = ("date", "diarienummer") if "blocks" in _CONTENT_KEYS[fmt] else ()
    fields = _fields(
        raw_fields, spec_name, (*_COMMON_KEYS, *_CONTENT_KEYS[fmt]), optional
    )
    file = _text(fields["file"], f"{spec_name}: file")
    if f"{file}.json" != spec_name:
        raise ValueError(f"{spec_name}: a spec for {file} must be named {file}.json")
    if Path(file).suffix != f".{fmt}":
        raise ValueError(f"{spec_name}: the file suffix must be .{fmt}")
    source = _fields(fields["source"], f"{spec_name}: source", ("catalogue_id", "url"))
    content: DocumentContent | tuple[Sheet, ...] | JsonData
    if "blocks" in fields:
        content = _document(fields, spec_name, fmt)
    elif "sheets" in fields:
        content = _sheets(fields["sheets"], spec_name, fmt)
    elif _finite(fields["data"]):
        content = JsonData(fields["data"])
    else:
        raise ValueError(f"{spec_name}: data holds a number that is not finite")
    return FixtureSpec(
        file=file,
        format=fmt,
        source=Source(
            _text(source["catalogue_id"], f"{spec_name}: source catalogue_id"),
            _text(source["url"], f"{spec_name}: source url"),
        ),
        description=_text(fields["description"], f"{spec_name}: description"),
        content=content,
    )


def cell_text(value: Cell) -> str:
    """A cell as a CSV export writes it: empty for null."""

    return "" if value is None else str(value)


def configure_document(
    document: DocxDocument, letterhead: Letterhead, title: str
) -> None:
    properties = document.core_properties
    properties.author = AUTHOR
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP
    properties.last_printed = FIXED_TIMESTAMP
    properties.title = title

    styles = document.styles
    normal = cast(ParagraphStyle, styles["Normal"])
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for name, size, bold in (
        ("Municipal Title", 20, True),
        ("Municipal Heading", 13, True),
        ("Municipal Metadata", 9, False),
    ):
        style = cast(
            ParagraphStyle,
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH),  # pyright: ignore[reportUnknownMemberType]  # python-docx leaves its parameters untyped
        )
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold

    for section in document.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.0)
        header = section.header.paragraphs[0]
        header.text = f"{letterhead.organisation.upper()}  |  {letterhead.unit.upper()}"
        header.style = "Municipal Metadata"
        footer = section.footer.paragraphs[0]
        footer.text = "  |  ".join(
            part for part in (letterhead.organisation, letterhead.address) if part
        )
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.style = "Municipal Metadata"


def _normalized_zip(raw: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(raw))
    normalized = io.BytesIO()
    with (
        source,
        zipfile.ZipFile(
            normalized,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as target,
    ):
        for member in sorted(source.infolist(), key=lambda item: item.filename):
            info = zipfile.ZipInfo(member.filename, FIXED_ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(member.filename))
    return normalized.getvalue()


def document_bytes(document: DocxDocument) -> bytes:
    raw = io.BytesIO()
    document.save(raw)
    return _normalized_zip(raw.getvalue())


def add_metadata_table(document: DocxDocument, rows: Sequence[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value


_STYLESHEET = """
@page {
  size: A4;
  margin: 26mm 20mm 22mm 24mm;
  @top-left { content: string(organisation) " · " string(unit); font: 7.5pt Arial, "Liberation Sans", "DejaVu Sans",
    sans-serif; color: #555; }
  @top-right { content: "Sida " counter(page) " av " counter(pages); font: 7.5pt Arial, "Liberation Sans",
    "DejaVu Sans", sans-serif; color: #555; }
}
@page :first { @top-left { content: none; } }
body { font: 10pt/1.4 Arial, "Liberation Sans", "DejaVu Sans", sans-serif; color: #111;
  font-variant-ligatures: none; }
.letterhead { border-bottom: 1.2pt solid #1f4e79; padding-bottom: 6pt; margin-bottom: 16pt; }
.letterhead p { margin: 0; }
.organisation { font-size: 15pt; font-weight: bold; color: #1f4e79; string-set: organisation content(); }
.unit { string-set: unit content(); }
.address { font-size: 8pt; color: #444; }
.metadata { margin-top: 4pt; text-align: right; font-size: 9pt; }
h1 { font-size: 15pt; margin: 0 0 10pt; }
h2 { font-size: 11.5pt; margin: 14pt 0 4pt; page-break-after: avoid; }
p { margin: 0 0 7pt; }
ul { margin: 0 0 8pt; padding-left: 16pt; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 10pt; font-size: 9pt; }
th, td { border: 0.6pt solid #666; padding: 2.5pt 4pt; text-align: left; vertical-align: top; }
th { background: #e8eef5; }
tr { page-break-inside: avoid; }
.signature { margin-top: 28pt; width: 70mm; border-top: 0.6pt solid #111; padding-top: 3pt; }
.signature p { margin: 0; }
.page-break { page-break-before: always; }
"""


def _html_block(block: Block) -> str:
    match block:
        case Heading(text=text):
            return f"<h2>{html.escape(text)}</h2>"
        case Paragraph(text=text):
            return f"<p>{html.escape(text)}</p>"
        case BulletList(items=items):
            return (
                "<ul>"
                + "".join(f"<li>{html.escape(item)}</li>" for item in items)
                + "</ul>"
            )
        case Table(columns=columns, rows=rows):
            head = "".join(f"<th>{html.escape(column)}</th>" for column in columns)
            body = "".join(
                "<tr>"
                + "".join(f"<td>{html.escape(cell)}</td>" for cell in row)
                + "</tr>"
                for row in rows
            )
            return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"
        case Signature(name=name, role=role):
            return (
                f'<div class="signature"><p>{html.escape(name)}</p>'
                f"<p>{html.escape(role)}</p></div>"
            )
        case PageBreak():
            return '<div class="page-break"></div>'
        case Control():
            return ""  # parse_spec keeps controls in docx


def _pdf(document: DocumentContent) -> bytes:
    letterhead = document.letterhead
    pinned = FIXED_TIMESTAMP.date().isoformat()
    metadata = "".join(
        f"<p>{label}: {html.escape(value)}</p>" for label, value in document.metadata()
    )
    address = (
        f'<p class="address">{html.escape(letterhead.address)}</p>'
        if letterhead.address
        else ""
    )
    page = (
        f'<!doctype html><html lang="sv"><head><meta charset="utf-8">'
        f"<title>{html.escape(document.title)}</title>"
        f'<meta name="author" content="{AUTHOR}">'
        f'<meta name="dcterms.created" content="{pinned}">'
        f'<meta name="dcterms.modified" content="{pinned}">'
        f"<style>{_STYLESHEET}</style></head><body>"
        f'<header class="letterhead"><p class="organisation">{html.escape(letterhead.organisation)}</p>'
        f'<p class="unit">{html.escape(letterhead.unit)}</p>{address}'
        f'<div class="metadata">{metadata}</div></header>'
        f"<h1>{html.escape(document.title)}</h1>"
        + "".join(_html_block(block) for block in document.blocks)
        + "</body></html>"
    )
    # Without a target, write_pdf returns the bytes.
    return cast(bytes, weasyprint.HTML(string=page).write_pdf())  # pyright: ignore[reportUnknownMemberType]  # WeasyPrint publishes no stubs


def _docx(document: DocumentContent) -> bytes:
    output = docx.Document()
    configure_document(output, document.letterhead, document.title)
    metadata = document.metadata()
    if metadata:
        add_metadata_table(output, metadata)
    title = output.add_paragraph(document.title, style="Municipal Title")
    title.paragraph_format.space_before = Pt(12)
    for block in document.blocks:
        match block:
            case Heading(text=text):
                output.add_paragraph(text, style="Municipal Heading")
            case Paragraph(text=text):
                output.add_paragraph(text)
            case BulletList(items=items):
                for item in items:
                    output.add_paragraph(item, style="List Bullet")
            case Table(columns=columns, rows=rows):
                table = output.add_table(rows=0, cols=len(columns))
                table.style = "Table Grid"
                for row in (columns, *rows):
                    for cell, text in zip(table.add_row().cells, row, strict=True):
                        cell.text = text
                for cell in table.rows[0].cells:
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
            case Signature(name=name, role=role):
                paragraph = output.add_paragraph(name)
                paragraph.paragraph_format.space_before = Pt(18)
                paragraph.add_run().add_break()
                paragraph.add_run(role)
            case PageBreak():
                output.add_page_break()
            case Control(kind="rich", tag=tag, label=label, hint=hint):
                append_rich_control(output, tag=tag, label=label, hint=hint)
            case Control(tag=tag, label=label, hint=hint):
                append_text_control(
                    output.add_paragraph(), tag=tag, label=label, hint=hint
                )
    # The product's helpers number controls across the process; a fixture numbers its own so it renders the same.
    for number, identifier in enumerate(output.element.iter(qn("w:id")), start=1):
        identifier.set(qn("w:val"), str(number))
    return document_bytes(output)


def _txt(document: DocumentContent) -> bytes:
    letterhead = document.letterhead
    parts = [
        "\n".join(
            part
            for part in (letterhead.organisation, letterhead.unit, letterhead.address)
            if part
        )
    ]
    metadata = document.metadata()
    if metadata:
        parts.append("\n".join(f"{label}: {value}" for label, value in metadata))
    parts.append(document.title)
    for block in document.blocks:
        match block:
            case Heading(text=text) | Paragraph(text=text):
                parts.append(text)
            case BulletList(items=items):
                parts.append("\n".join(f"- {item}" for item in items))
            case Table(columns=columns, rows=rows):
                parts.append("\n".join(" | ".join(row) for row in (columns, *rows)))
            case Signature(name=name, role=role):
                parts.append(f"{name}\n{role}")
            case PageBreak() | Control():
                pass  # parse_spec keeps page breaks out of txt and controls in docx
    return ("\n\n".join(parts) + "\n").encode("utf-8")


def _csv(sheet: Sheet) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(sheet.columns)
    writer.writerows([cell_text(value) for value in row] for row in sheet.rows)
    return buffer.getvalue().encode("utf-8")


def _xlsx(sheets: Sequence[Sheet]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for sheet in sheets:
        worksheet = workbook.create_sheet(sheet.name)
        worksheet.append(sheet.columns)
        for row in sheet.rows:
            worksheet.append(row)
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for index in range(len(sheet.columns)):
            width = max(
                len(cell_text(row[index])) for row in (sheet.columns, *sheet.rows)
            )
            worksheet.column_dimensions[get_column_letter(index + 1)].width = width + 2
    workbook.properties.creator = AUTHOR
    workbook.properties.created = FIXED_TIMESTAMP
    workbook.properties.modified = FIXED_TIMESTAMP
    raw = io.BytesIO()
    # ExcelWriter directly: openpyxl's save() stamps the current time as the modified date.
    with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED) as archive:
        ExcelWriter(workbook, archive).save()
    return _normalized_zip(raw.getvalue())


_DOCUMENT_RENDERERS: dict[str, Callable[[DocumentContent], bytes]] = {
    "pdf": _pdf,
    "docx": _docx,
    "txt": _txt,
}


def render(spec: FixtureSpec) -> bytes:
    match spec.content:
        case DocumentContent() as document:
            return _DOCUMENT_RENDERERS[spec.format](document)
        case JsonData(value=value):
            return (
                json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
            ).encode("utf-8")
        case sheets:
            return _csv(sheets[0]) if spec.format == "csv" else _xlsx(sheets)
