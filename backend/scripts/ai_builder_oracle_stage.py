#!/usr/bin/env python3
"""Stage the request material an expert gold-spec author is allowed to see.

An author of arm O gets what a skilled flow author gets from a requester: the
request text, the files attached to it, the requester's answers to intake
questions and the platform's authoring documentation. It never gets the case's
oracle. The cases are read by the harness's own parser, so the answers staged
are the RESOLVED ones the Builder arm is sent (the synthetic profile's answers
with the case's own on top, `_merged_configured_answers`), and only four
fields of a parsed case are projected (PROJECTED_FIELDS): everything else
(`expected`, `execution` with its runtime files, checkpoints and oracle,
`edit`, `note`, cohorts, the source URL) is never put in the staged tree, which
carries a manifest of every file it holds.

The staging is deterministic, and `verify_staged_case` recomputes everything a
freeze relies on from the staged BYTES: a manifest is never trusted for a
digest it claims.

usage: ai_builder_oracle_stage.py --cases-file CASES --selection SELECTION.json --out DIR
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib
import io
import json
import sys
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any, cast

_SCRIPTS = Path(__file__).resolve().parent
# Standalone script: a test loads it by path, which does not put this directory
# on the import path by itself.
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))

from ai_builder_intake_answers import (  # noqa: E402
    intake_sha256,
    render_answers,
    staged_intake_sha256,
)

_BACKEND = _SCRIPTS.parent
_REPO = _BACKEND.parent
FIXTURES = _SCRIPTS / "fixtures" / "ai_builder_battle"

# The only fields of a parsed case that reach the staged tree. `answers` is the
# case's resolved `configured_question_answers`.
PROJECTED_FIELDS = ("id", "prompt", "attachments", "answers")

# The authoring documentation: platform docs and the spec schema, with the
# Builder's own sections left out (an author authors a spec, it does not drive
# a Builder). (repo-relative path, first line, last line); None = whole file.
AUTHOR_DOCS: tuple[tuple[str, int | None, int | None], ...] = (
    ("docs/flows/flow-developer-quickstart.md", 14, 280),
    ("docs/flows/flow-developer-quickstart.md", 486, 559),
    ("docs/flows/flow-developer-quickstart.md", 589, 648),
    (
        "frontend/apps/docs-site/src/content/guides/flows/designing-flows.mdx",
        None,
        None,
    ),
    ("frontend/apps/docs-site/src/content/docs/flows-best-practices.mdx", None, None),
    ("frontend/apps/docs-site/src/content/docs/flows-long-material.mdx", None, None),
    ("backend/src/eneo/flows/flow_authoring_spec.py", None, None),
    ("backend/src/eneo/flows/flow_capability_manifest.py", None, None),
)


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


class MaterialError(ValueError):
    """A staged tree is not what its manifest says, or not what the Builder gets."""


def request_material(case: Any) -> dict[str, Any]:
    """The projection of a parsed case; nothing else leaves this function."""

    return {
        "id": case.case_id,
        "prompt": case.prompt,
        "attachments": list(case.attachments),
        "answers": dict(case.configured_question_answers or {}),
    }


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if suffix == ".docx":
        from eneo.flows.runtime.docx_template_runtime import extract_docx_text

        return extract_docx_text(path.read_bytes())
    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), data_only=True)
        out: list[str] = []
        for sheet in workbook.worksheets:
            out.append(f"## Blad: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                out.append(";".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(out)
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(path.read_text(encoding="utf-8")))
        return "\n".join(";".join(row) for row in rows)
    return path.read_text(encoding="utf-8")


def stage_case(
    material: Mapping[str, Any], out: Path, *, fixtures: Path = FIXTURES
) -> dict[str, Any]:
    case_dir = out / str(material["id"])
    (case_dir / "attachments").mkdir(parents=True, exist_ok=True)
    prompt = cast(str, material["prompt"])
    answers = cast(Mapping[str, Any], material.get("answers") or {})
    # The request, then the answers (empty when there are none): joined by
    # `join_intake`, the two files are exactly the text the Builder arm
    # receives up front.
    files: dict[str, bytes] = {
        "request.md": prompt.encode("utf-8"),
        "answers.md": (render_answers(answers) if answers else "").encode("utf-8"),
    }
    for name in cast(Sequence[str], material.get("attachments") or []):
        source = fixtures / name
        files[f"attachments/{name}"] = source.read_bytes()
        files[f"attachments/{name}.txt"] = _extract_text(source).encode("utf-8")
    for relative, payload in files.items():
        (case_dir / relative).write_bytes(payload)
    manifest = {
        "case_id": material["id"],
        "projected_fields": list(PROJECTED_FIELDS),
        "intake_message_sha256": intake_sha256(prompt, answers),
        "files": {
            relative: _sha256(payload) for relative, payload in sorted(files.items())
        },
    }
    (case_dir / "MATERIAL.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=1, sort_keys=True),
        encoding="utf-8",
    )
    return manifest


def verify_staged_case(case_dir: Path) -> str:
    """The intake digest of a staged case, from its bytes; refuses any drift.

    Every file on disk must be listed in `MATERIAL.json` with its hash, no
    listed file may be missing, and if the manifest states an intake digest it
    must be the one the staged `request.md` and `answers.md` give. The value
    returned is always the recomputed one.
    """

    manifest = json.loads((case_dir / "MATERIAL.json").read_text(encoding="utf-8"))
    if manifest.get("projected_fields") != list(PROJECTED_FIELDS):
        raise MaterialError(f"{case_dir.name}: MATERIAL.json projects other fields.")
    listed = cast(Mapping[str, str], manifest.get("files") or {})
    on_disk = {
        str(path.relative_to(case_dir)): _sha256(path.read_bytes())
        for path in sorted(case_dir.rglob("*"))
        if path.is_file() and path.name != "MATERIAL.json"
    }
    if on_disk != dict(listed):
        changed = sorted(
            name
            for name in set(on_disk) | set(listed)
            if on_disk.get(name) != listed.get(name)
        )
        raise MaterialError(
            f"{case_dir.name}: staged files differ from MATERIAL.json: {changed}."
        )
    recomputed = staged_intake_sha256(
        (case_dir / "request.md").read_bytes(), (case_dir / "answers.md").read_bytes()
    )
    claimed = manifest.get("intake_message_sha256")
    if claimed is not None and claimed != recomputed:
        raise MaterialError(
            f"{case_dir.name}: MATERIAL.json claims an intake digest the staged "
            "brief does not have."
        )
    return recomputed


def verify_docs(material_dir: Path) -> str:
    """The digest of `docs/DOCS.json`, after checking every doc against it."""

    docs = material_dir / "docs"
    listed = json.loads((docs / "DOCS.json").read_text(encoding="utf-8"))
    on_disk = {
        path.name: _sha256(path.read_bytes())
        for path in sorted(docs.iterdir())
        if path.is_file() and path.name != "DOCS.json"
    }
    if on_disk != listed:
        raise MaterialError("the staged docs differ from DOCS.json.")
    return _sha256((docs / "DOCS.json").read_bytes())


def material_digest(material_dir: Path, case_ids: Sequence[str]) -> str:
    """One digest over the staged tree: each case's `MATERIAL.json` bytes, by id."""

    return _sha256(
        json.dumps(
            {
                case_id: _sha256(
                    (material_dir / case_id / "MATERIAL.json").read_bytes()
                )
                for case_id in sorted(case_ids)
            },
            sort_keys=True,
        ).encode("utf-8")
    )


def stage_docs(out: Path, *, repo: Path = _REPO) -> dict[str, str]:
    docs = out / "docs"
    docs.mkdir(parents=True, exist_ok=True)
    digests: dict[str, str] = {}
    for relative, first, last in AUTHOR_DOCS:
        lines = (repo / relative).read_text(encoding="utf-8").splitlines()
        chosen = lines[(first or 1) - 1 : last or len(lines)]
        name = relative.replace("/", "__") + (f".{first}-{last}" if first else "")
        payload = ("\n".join(chosen) + "\n").encode("utf-8")
        (docs / name).write_bytes(payload)
        digests[name] = _sha256(payload)
    (docs / "DOCS.json").write_text(
        json.dumps(digests, indent=1, sort_keys=True), encoding="utf-8"
    )
    return digests


def _tree_files(root: Path) -> dict[str, bytes]:
    return {
        str(path.relative_to(root)): path.read_bytes()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def verify_against_source(
    material_dir: Path,
    case_ids: Sequence[str],
    *,
    cases_file: Path,
    fixtures: Path,
    repo: Path,
) -> None:
    """The staged material is what staging its SOURCES produces, byte for byte.

    `verify_staged_case` and `verify_docs` prove a staged tree is consistent
    with the manifests it carries; whoever can edit a file can edit the manifest
    beside it. So the whole material (every selected case's request, answers,
    attachments, attachment renderings and manifest, and the documentation
    slices) is regenerated here from the corpus, the fixtures and the docs of
    the frozen tree, into a scratch directory, and compared with what the
    author was given. A missing, extra or different file refuses the freeze.
    """

    import tempfile

    cases = load_cases(cases_file)
    unknown = sorted(set(case_ids) - set(cases))
    if unknown:
        raise MaterialError(f"cases the frozen corpus does not hold: {unknown}.")
    with tempfile.TemporaryDirectory() as scratch:
        expected_root = Path(scratch)
        for case_id in case_ids:
            stage_case(
                request_material(cases[case_id]), expected_root, fixtures=fixtures
            )
        stage_docs(expected_root, repo=repo)
        expected = _tree_files(expected_root)
    staged = _tree_files(material_dir)
    differing = sorted(
        name
        for name in set(expected) | set(staged)
        if expected.get(name) != staged.get(name)
    )
    if differing:
        raise MaterialError(
            "the staged material is not what the frozen corpus, fixtures and docs "
            f"produce; differing files: {differing[:12]}"
            + (f" (and {len(differing) - 12} more)" if len(differing) > 12 else "")
        )


def _harness() -> Any:
    """The harness module, loaded as `run_harness.py` loads it (tables first)."""

    name = "ai_builder_api_battle_test"
    if name in sys.modules:
        return sys.modules[name]
    importlib.import_module("eneo.database.tables")
    return importlib.import_module(name)


def load_cases(cases_file: Path) -> dict[str, Any]:
    """The corpus as the harness parses it: the answers already resolved."""

    return {case.case_id: case for case in _harness()._read_cases_file(cases_file)}


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cases-file", required=True)
    parser.add_argument("--selection", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args(argv)
    cases = load_cases(Path(args.cases_file))
    selection = json.loads(Path(args.selection).read_text())
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=False)
    for entry in selection["cases"]:
        manifest = stage_case(request_material(cases[entry["id"]]), out)
        print(f"{entry['id']}: {len(manifest['files'])} files")
    docs = stage_docs(out)
    print(f"docs: {len(docs)} files")
    return 0


if __name__ == "__main__":
    sys.exit(main())
