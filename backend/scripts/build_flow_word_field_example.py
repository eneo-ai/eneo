"""Build the downloadable example for filling separate places in a Word template."""

from pathlib import Path
from typing import cast

from docx import Document
from docx.oxml.document import CT_Document
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from lxml import etree

from eneo.flows.runtime.document_rendering.docx_content_controls import (
    append_rich_control,
    append_text_control,
)

_W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_OUTPUT = (
    Path(__file__).resolve().parents[2]
    / "frontend/apps/web/static/examples/eneo-word-template-fields.docx"
)


def build_example() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    style_sizes = [("Normal", 11), ("Title", 24)] + [
        (f"Heading {level}", 14 if level == 1 else 12 if level == 2 else 11)
        for level in range(1, 10)
    ]
    for name, size in style_sizes:
        style = document.styles[name]
        assert isinstance(style, ParagraphStyle)
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
    heading = document.styles["Heading 1"]
    normal = document.styles["Normal"]
    assert isinstance(heading, ParagraphStyle) and isinstance(normal, ParagraphStyle)
    heading.paragraph_format.space_before = Pt(16)
    normal.paragraph_format.line_spacing = 1.15
    run_properties = normal.element.find(qn("w:rPr"))
    assert run_properties is not None
    language = etree.SubElement(run_properties, qn("w:lang"))
    language.set(qn("w:val"), "sv-SE")
    for setting in document.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            parent = border.getparent()
            assert parent is not None
            parent.remove(border)
    document.core_properties.title = "Utredningsunderlag"
    document.core_properties.subject = "Exempelmall med fyra namngivna fält för Eneo"
    document.core_properties.language = "sv-SE"
    document.add_heading("Utredningsunderlag", level=0)
    name = document.add_paragraph()
    name.add_run("Namn: ").bold = True
    append_text_control(
        name,
        tag="namn",
        label="Namn",
        hint="Koppla till ett namnfält i Eneo",
    )
    sections = (
        (
            "bakgrund",
            "Bakgrund",
            "Här fylls bakgrundstexten i. Koppla fältet Bakgrund till text från ett steg i Eneo.",
        ),
        (
            "bedomning",
            "Bedömning",
            "Här fylls bedömningen i. Välj den text som ska hamna under denna rubrik.",
        ),
        (
            "nasta_steg",
            "Nästa steg",
            "Här fylls nästa steg i. Välj en textkälla, eller Lämna tomt om avsnittet ska vara tomt.",
        ),
    )
    for tag, label, hint in sections:
        document.add_heading(label, level=1)
        control = append_rich_control(document, tag=tag, label=label, hint=hint)
        content = cast(BaseOxmlElement, control).find(qn("w:sdtContent"))
        assert content is not None
        paragraph = etree.SubElement(content, qn("w:p"))
        run = etree.SubElement(paragraph, qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = (
            "Hjälptexten ersätts när flödet körs. Rubriken ovanför ligger utanför "
            "fältet och finns kvar. Klicka i fältet i Word för att se dess namn."
        )

    # These are Word's editing boundaries, not printed borders or text colors.
    root = cast(CT_Document, document.element)
    for control in root.body.iter(qn("w:sdt")):
        properties = control.find(qn("w:sdtPr"))
        assert properties is not None
        appearance = etree.SubElement(
            properties, f"{{{_W15}}}appearance", nsmap={"w15": _W15}
        )
        appearance.set(f"{{{_W15}}}val", "boundingBox")
        color = etree.SubElement(properties, f"{{{_W15}}}color")
        color.set(qn("w:val"), "2343DA")
    etree.cleanup_namespaces(
        document.element,
        top_nsmap={"w15": _W15},
        keep_ns_prefixes=[
            prefix for prefix in document.element.nsmap if prefix is not None
        ],
    )
    ignorable = document.element.get(f"{{{_MC}}}Ignorable", "").split()
    document.element.set(f"{{{_MC}}}Ignorable", " ".join([*ignorable, "w15"]))
    document.save(str(_OUTPUT))


if __name__ == "__main__":
    build_example()
