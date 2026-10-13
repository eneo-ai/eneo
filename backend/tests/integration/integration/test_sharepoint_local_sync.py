"""Full SharePoint syncs over a local folder, end to end through the publish path.

The Graph client is replaced by a directory walk; everything after the bytes
arrive is production code: byte retention for tabular files, versioning of
originals, and the quota fallback.
"""

from __future__ import annotations

import shutil
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import UUID, uuid4

import pytest
import sqlalchemy as sa
from dependency_injector import providers
from openpyxl import load_workbook

from eneo.database.tables.ai_models_table import EmbeddingModels
from eneo.database.tables.info_blob_chunk_table import InfoBlobChunks
from eneo.database.tables.info_blobs_table import InfoBlobs, InfoBlobVersionState
from eneo.database.tables.object_content_table import ObjectContents
from eneo.database.tables.spaces_table import Spaces
from eneo.database.tables.tenant_table import Tenants
from eneo.files.chunk_embedding_list import ChunkEmbeddingList
from eneo.info_blobs.info_blob_repo import InfoBlobRepository
from eneo.integration.domain.entities.integration import Integration
from eneo.integration.domain.entities.integration_knowledge import (
    IntegrationKnowledge,
)
from eneo.integration.domain.entities.tenant_integration import TenantIntegration
from eneo.integration.domain.entities.user_integration import UserIntegration
from eneo.integration.infrastructure.content_service.sharepoint_content_service import (
    SharePointContentService,
)
from eneo.object_content.content import ContentState
from tests.integration.fixtures.sharepoint_local_client import (
    FIXTURE_DRIVE_ID,
    FIXTURE_ROOT_ID,
    FIXTURE_WEB_ORIGIN,
    LocalFolderSharePointClient,
)

FIXTURE_FILES = Path(__file__).resolve().parent.parent / "fixtures" / "files"
CLIENT_PATH = (
    "eneo.integration.infrastructure.content_service."
    "sharepoint_content_service.SharePointContentClient"
)


def _embedding_result(*, model, chunks):
    result = ChunkEmbeddingList()
    result.add(chunks, [[0.7, 0.8, 0.9] for _ in chunks])
    return result


async def _seed_source(container) -> UUID:
    session = container.session()
    user = container.user()
    space = await session.scalar(
        sa.select(Spaces).where(
            Spaces.tenant_id == user.tenant_id, Spaces.user_id == user.id
        )
    )
    if space is None:
        space = Spaces(
            name=f"local sync {uuid4().hex}", tenant_id=user.tenant_id, user_id=user.id
        )
        session.add(space)
        await session.flush()
    model_row = (await session.scalars(sa.select(EmbeddingModels).limit(1))).one()
    embedding_model = await container.embedding_model_repo2().one(model_id=model_row.id)

    integration = await container.integration_repo().add(
        Integration(
            name="SharePoint", description="fixture", integration_type="sharepoint"
        )
    )
    tenant_integration = await container.tenant_integration_repo().add(
        TenantIntegration(tenant_id=user.tenant_id, integration=integration)
    )
    user_integration = await container.user_integration_repo().add(
        UserIntegration(
            tenant_integration=tenant_integration, user_id=user.id, authenticated=True
        )
    )
    knowledge = await container.integration_knowledge_repo().add(
        obj=IntegrationKnowledge(
            name="Local folder",
            user_integration=user_integration,
            embedding_model=embedding_model,
            tenant_id=user.tenant_id,
            space_id=space.id,
            url=FIXTURE_WEB_ORIGIN,
            drive_id=FIXTURE_DRIVE_ID,
            folder_id=FIXTURE_ROOT_ID,
            resource_type="onedrive",
            selected_item_type="folder",
        )
    )
    return knowledge.id


async def _sync(container, monkeypatch, *, knowledge_id: UUID, folder: Path):
    embeddings = AsyncMock()
    embeddings.get_embeddings.side_effect = _embedding_result
    container.create_embeddings_service.override(providers.Object(embeddings))
    monkeypatch.setattr(
        CLIENT_PATH, lambda **kwargs: LocalFolderSharePointClient(root=folder, **kwargs)
    )
    service = container.sharepoint_content_service()
    service.oauth_token_repo.one = AsyncMock(  # type: ignore[method-assign]
        return_value=SimpleNamespace(
            id=uuid4(),
            access_token="fixture",
            base_url=FIXTURE_WEB_ORIGIN,
            is_sharepoint=True,
        )
    )
    await service.pull_content(
        token_id=uuid4(),
        integration_knowledge_id=knowledge_id,
        drive_id=FIXTURE_DRIVE_ID,
        folder_id=FIXTURE_ROOT_ID,
        resource_type="onedrive",
    )
    # What the sync recorded about itself: counts and per-file notes.
    (log,) = await container.sync_log_repo().get_by_integration_knowledge(
        knowledge_id, tenant_id=container.user().tenant_id, limit=1
    )
    return log.metadata or {}


async def _documents(session, knowledge_id: UUID) -> dict[str, list[InfoBlobs]]:
    rows = (
        await session.scalars(
            sa.select(InfoBlobs)
            .where(InfoBlobs.integration_knowledge_id == knowledge_id)
            .order_by(InfoBlobs.created_at, InfoBlobs.id)
        )
    ).all()
    by_title: dict[str, list[InfoBlobs]] = {}
    for row in rows:
        by_title.setdefault(row.title or "", []).append(row)
    return by_title


async def _active(session, knowledge_id: UUID, title: str) -> InfoBlobs:
    versions = (await _documents(session, knowledge_id))[title]
    (active,) = [
        v for v in versions if v.version_state == InfoBlobVersionState.ACTIVE.value
    ]
    return active


async def _chunk_count(session, blob_id: UUID) -> int:
    return (
        await session.scalar(
            sa.select(sa.func.count()).where(InfoBlobChunks.info_blob_id == blob_id)
        )
        or 0
    )


def _append_row(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    workbook["Sales"].append(["2026-04", "North", 200])
    workbook.save(workbook_path)


BIG_CSV_ROWS = 20_000


def _write_big_csv(path: Path) -> None:
    """A delimited table well over the index budget: indexed as a sample."""
    path.write_text(
        "id,name\n" + "".join(f"{i},person-{i}\n" for i in range(BIG_CSV_ROWS))
    )


def _edit_row(path: Path, index: int) -> None:
    """Change one row beyond the sample without changing the file's size."""
    text = path.read_text()
    edited = text.replace(f"{index},person-{index}\n", f"{index},edited-{index}\n")
    assert edited != text and len(edited) == len(text)
    path.write_text(edited)


@pytest.mark.integration
async def test_tabular_originals_follow_the_document_through_syncs(
    db_container, tmp_path, monkeypatch
):
    folder = tmp_path / "drive"
    folder.mkdir()
    shutil.copy(FIXTURE_FILES / "sales.xlsx", folder / "sales.xlsx")
    (folder / "notes.txt").write_text("Minutes from the budget meeting.\n")

    async with db_container() as container:
        knowledge_id = await _seed_source(container)

    # First sync: the spreadsheet is indexed and keeps its bytes; prose keeps none.
    async with db_container() as container:
        stats = await _sync(
            container, monkeypatch, knowledge_id=knowledge_id, folder=folder
        )
        assert stats["files_processed"] == 2, stats
        session = container.session()
        repo = InfoBlobRepository(session)
        sheet = await _active(session, knowledge_id, "sales.xlsx")
        notes = await _active(session, knowledge_id, "notes.txt")
        original = await repo.get_original(sheet.id)
        assert original is not None and original.usable
        assert original.original_filename == "sales.xlsx"
        assert original.media_type == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert await _chunk_count(session, sheet.id) > 0
        assert await repo.get_original(notes.id) is None
        first_sheet_id, first_content_id = sheet.id, original.content_id

    # Unchanged files: no new version, no new bytes.
    async with db_container() as container:
        await _sync(container, monkeypatch, knowledge_id=knowledge_id, folder=folder)
        session = container.session()
        assert len((await _documents(session, knowledge_id))["sales.xlsx"]) == 1
        original = await InfoBlobRepository(session).get_original(first_sheet_id)
        assert original is not None and original.content_id == first_content_id

    # A changed workbook: a new active version holds the new bytes, the
    # superseded version keeps its text but releases its bytes.
    _append_row(folder / "sales.xlsx")
    async with db_container() as container:
        await _sync(container, monkeypatch, knowledge_id=knowledge_id, folder=folder)
        session = container.session()
        repo = InfoBlobRepository(session)
        versions = (await _documents(session, knowledge_id))["sales.xlsx"]
        assert [v.version_state for v in versions] == ["superseded", "active"]
        assert versions[0].id == first_sheet_id
        assert await repo.get_original(first_sheet_id) is None
        current = await repo.get_original(versions[1].id)
        assert current is not None and current.content_id != first_content_id
        state, reference_count = (
            await session.execute(
                sa.select(ObjectContents.state, ObjectContents.reference_count).where(
                    ObjectContents.id == first_content_id
                )
            )
        ).one()
        assert (state, reference_count) == (ContentState.DELETE_PENDING.value, 0)
        assert versions[0].text  # text history stays


@pytest.mark.integration
async def test_original_is_dropped_not_the_document_when_quota_is_full(
    db_container, tmp_path, monkeypatch
):
    folder = tmp_path / "drive"
    folder.mkdir()
    shutil.copy(FIXTURE_FILES / "sales.xlsx", folder / "sales.xlsx")

    async with db_container() as container:
        knowledge_id = await _seed_source(container)
        session = container.session()
        user = container.user()
        retained = await InfoBlobRepository(session).get_retained_size_of_tenant(
            user.tenant_id
        )
        # Room for the extracted text, not for the workbook's bytes.
        await session.execute(
            sa.update(Tenants)
            .where(Tenants.id == user.tenant_id)
            .values(quota_limit=retained + 2_000)
        )

    async with db_container() as container:
        stats = await _sync(
            container, monkeypatch, knowledge_id=knowledge_id, folder=folder
        )
        assert stats["files_processed"] == 1, stats
        assert any(
            detail["file"] == "sales.xlsx"
            and detail["reason"].startswith("Original not kept: quota")
            for detail in stats["skipped_details"]
        ), stats["skipped_details"]
        session = container.session()
        sheet = await _active(session, knowledge_id, "sales.xlsx")
        # The bytes were dropped, the rows are still indexed.
        assert await InfoBlobRepository(session).get_original(sheet.id) is None
        assert await _chunk_count(session, sheet.id) > 0
        assert "North" in (sheet.text or "")


@pytest.mark.integration
async def test_same_size_edit_near_the_quota_keeps_its_new_original(
    db_container, tmp_path, monkeypatch
):
    """Replacing an original is charged for its growth, not as a second copy.

    An edit beyond the indexed sample leaves the text unchanged; the bytes are
    what changed. Near the quota the sync must still swap the bytes, or the
    table tools keep computing over the old file.
    """
    folder = tmp_path / "drive"
    folder.mkdir()
    _write_big_csv(folder / "big.csv")

    async with db_container() as container:
        knowledge_id = await _seed_source(container)

    async with db_container() as container:
        await _sync(container, monkeypatch, knowledge_id=knowledge_id, folder=folder)
        session = container.session()
        repo = InfoBlobRepository(session)
        first = await _active(session, knowledge_id, "big.csv")
        assert (first.text or "").startswith(
            f"File: big.csv\n\nLarge table: {BIG_CSV_ROWS} rows."
        )
        first_original = await repo.get_original(first.id)
        assert first_original is not None
        user = container.user()
        retained = await repo.get_retained_size_of_tenant(user.tenant_id)
        # Room for the re-indexed sample and its chunks; not for a second copy
        # of the bytes.
        await session.execute(
            sa.update(Tenants)
            .where(Tenants.id == user.tenant_id)
            .values(quota_limit=retained + 20_000)
        )

    _edit_row(folder / "big.csv", 10_000)
    async with db_container() as container:
        stats = await _sync(
            container, monkeypatch, knowledge_id=knowledge_id, folder=folder
        )
        assert not any(
            detail["reason"].startswith("Original not kept")
            for detail in stats["skipped_details"]
        ), stats["skipped_details"]
        session = container.session()
        repo = InfoBlobRepository(session)
        versions = (await _documents(session, knowledge_id))["big.csv"]
        assert [v.version_state for v in versions] == ["superseded", "active"]
        assert await repo.get_original(versions[0].id) is None
        current = await repo.get_original(versions[1].id)
        assert current is not None
        assert current.sha256 != first_original.sha256


@pytest.mark.integration
async def test_large_table_whose_bytes_are_dropped_is_indexed_in_full(
    db_container, tmp_path, monkeypatch
):
    """The sample promises the full file to the table tools; without the
    original, every row is indexed instead and the old bytes are released."""
    folder = tmp_path / "drive"
    folder.mkdir()
    _write_big_csv(folder / "big.csv")

    async with db_container() as container:
        knowledge_id = await _seed_source(container)

    async with db_container() as container:
        await _sync(container, monkeypatch, knowledge_id=knowledge_id, folder=folder)
        session = container.session()
        first = await _active(session, knowledge_id, "big.csv")
        assert await InfoBlobRepository(session).get_original(first.id) is not None

    # The deployment's knowledge-file limit now sits below the file.
    real_admission = SharePointContentService._upload_admission

    async def small_limit(self):
        snapshot = await real_admission(self)
        return replace(snapshot, knowledge_file_maximum_bytes=1_000)

    monkeypatch.setattr(SharePointContentService, "_upload_admission", small_limit)

    _edit_row(folder / "big.csv", 10_000)
    async with db_container() as container:
        stats = await _sync(
            container, monkeypatch, knowledge_id=knowledge_id, folder=folder
        )
        assert any(
            detail["file"] == "big.csv"
            and detail["reason"].startswith(
                "Original not kept: exceeds the knowledge file size limit"
            )
            for detail in stats["skipped_details"]
        ), stats["skipped_details"]
        session = container.session()
        repo = InfoBlobRepository(session)
        versions = (await _documents(session, knowledge_id))["big.csv"]
        assert [v.version_state for v in versions] == ["superseded", "active"]
        current = versions[1]
        assert "10000,edited-10000" in (current.text or "")
        assert not (current.text or "").startswith("File: big.csv")
        assert await repo.get_original(current.id) is None
        assert await repo.get_original(versions[0].id) is None
