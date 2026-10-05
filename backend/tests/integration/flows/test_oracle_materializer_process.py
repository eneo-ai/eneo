"""The oracle arm's materializer runs in a process that has no API lifespan.

The Container hands out the durable object-content runtime and every service
that reads or writes a File needs the database session manager, and both are
process singletons the API's lifespan starts. A harness process never ran that
lifespan, so the materializer brings both up itself, through the lifespan's own
start function, for the length of one apply and stops both after.

Each case runs the arm's default materializer (the one the harness builds from
the API key's user) in a fresh interpreter, the shape of the harness process,
against the integration database, and reads the flow back with plain SQL. A
process of its own is the point: this test process already owns both
singletons and an API-started one would hide the failure.
"""

from __future__ import annotations

import io
import json
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
import sqlalchemy as sa
from docx import Document

from eneo.database.database import sessionmanager
from eneo.database.tables.flow_tables import Flows, FlowSteps, FlowTemplateAssets
from eneo.database.tables.spaces_table import Spaces
from eneo.files.file_models import FileContentVariant, FileType
from eneo.files.file_protocol import PendingFileContent, PreparedFileUpload
from eneo.flows.flow_authoring_spec import (
    AssistantSpec,
    FlowDraftSpecCore,
    FormFieldSpec,
    InputSource,
    InputType,
    OutputMode,
    OutputType,
    StepSpec,
)
from eneo.flows.runtime.document_rendering.docx_content_controls import (
    append_text_control,
)
from tests.integration.flows.conftest import _flow_worker_environment

pytestmark = [pytest.mark.integration, pytest.mark.asyncio]

_BACKEND = Path(__file__).resolve().parents[3]
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# What a harness does with the arm: build the default materializer for the API
# key's user (the fake harness answers `/users/me/`), apply the frozen spec.
# The second apply proves the first one left nothing running behind it.
_HARNESS_PROCESS = """
import json
import sys
import time
from types import SimpleNamespace
from uuid import UUID

sys.path.insert(0, sys.argv[1])
import ai_builder_oracle_arm as arm
from eneo.database.database import sessionmanager
from eneo.flows.flow_authoring_spec import FlowDraftSpecCore
from eneo.object_content.runtime import ObjectContentRuntime, object_content_runtime

payload = json.load(sys.stdin)
harness = SimpleNamespace(
    _request_json=lambda **_: {"id": payload["user_id"]},
    _required_string=lambda body, key: body[key],
)
materializer = arm._default_materializer(
    harness, SimpleNamespace(base_url="http://stack", api_key="k" * 16)
)
spec = FlowDraftSpecCore.model_validate(payload["spec"])
gold = arm.GoldSpec(
    "materializer-process", spec, "spec.json", "0" * 64, payload["template_name"], {}
)
request = arm.MaterializeRequest(
    space_id=UUID(payload["space_id"]),
    gold=gold,
    spec=spec,
    template_file_id=(
        UUID(payload["template_file_id"]) if payload["template_file_id"] else None
    ),
    wait_until=time.monotonic() + 600.0,
)
first_error = None
if payload["fail_first_start"]:
    # The store cannot be reached while the first apply starts persistence.
    original = ObjectContentRuntime.validate_configuration
    attempts = []

    async def unreachable_once(self):
        attempts.append(1)
        if len(attempts) == 1:
            raise OSError("the store is unreachable")
        return await original(self)

    ObjectContentRuntime.validate_configuration = unreachable_once
    try:
        materializer.materialize(request)
    except arm.MaterializeInfrastructureError as error:
        first_error = str(error)
flows = [materializer.materialize(request).flow_id for _ in range(2)]


def session_manager_is_stopped():
    try:
        sessionmanager.create_session()
    except Exception:
        return True
    return False


print(
    json.dumps(
        {
            "first_error": first_error,
            "flow_ids": flows,
            "object_content_enabled": object_content_runtime.enabled,
            "session_manager_stopped": session_manager_is_stopped(),
        }
    )
)
"""


def _spec(*, template: bool) -> FlowDraftSpecCore:
    final_step = StepSpec(
        plan_step_ref="step_b",
        name="Skriv beslutet",
        assistant_spec=AssistantSpec(instructions="Skriv beslutet."),
        input_source=InputSource.PREVIOUS_STEP,
        input_type=InputType.TEXT,
        output_mode=OutputMode.PASS_THROUGH,
        output_type=OutputType.TEXT,
    )
    if template:
        final_step = final_step.model_copy(
            update={
                "output_mode": OutputMode.TEMPLATE_FILL,
                "output_type": OutputType.DOCX,
                "output_config": {"bindings": {"case_id": "{{ flow_input.case_id }}"}},
            }
        )
    return FlowDraftSpecCore(
        flow_name=f"Beslut {uuid4().hex[:8]}",
        flow_description="Skriver ett beslut.",
        steps=[
            StepSpec(
                plan_step_ref="step_a",
                name="Läs ärendet",
                assistant_spec=AssistantSpec(instructions="Läs ärendet noga."),
                input_source=InputSource.FLOW_INPUT,
                input_type=InputType.TEXT,
                output_type=OutputType.TEXT,
            ),
            final_step,
        ],
        form_fields=[FormFieldSpec(name="case_id", type="text", label="Ärende")],
    )


def _template_bytes() -> bytes:
    document = Document()
    append_text_control(
        document.add_paragraph("Ärende: "),
        tag="case_id",
        label="case_id",
        hint="case_id",
    )
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


async def _one_chunk(payload: bytes):
    yield payload


async def _space_and_template(db_container) -> tuple[UUID, UUID, UUID]:
    async with db_container() as container:
        user = container.user()
        space = Spaces(
            tenant_id=user.tenant_id,
            user_id=user.id,
            name=f"materializer-process-{uuid4().hex}",
        )
        container.session().add(space)
        await container.session().flush()
        file = await container.file_service().save_prepared_file(
            PreparedFileUpload(
                name="mall.docx",
                file_type=FileType.DOCUMENT,
                display_media_type=_DOCX_MIME,
                contents=(
                    PendingFileContent(
                        variant=FileContentVariant.ORIGINAL,
                        chunks=_one_chunk(_template_bytes()),
                        declared_media_type=_DOCX_MIME,
                        verified_media_type=_DOCX_MIME,
                    ),
                ),
            )
        )
        return user.id, space.id, file.id


# A process that starts persistence twice, then uses what the first start owns.
_SECOND_START_PROCESS = """
import asyncio
import json

from sqlalchemy import text

from eneo.database.database import sessionmanager
from eneo.object_content.runtime import object_content_runtime
from eneo.server.dependencies.lifespan import start_persistence, stop_persistence


async def main():
    await start_persistence()
    try:
        await start_persistence()
        refused = None
    except RuntimeError as error:
        refused = str(error)
    async with sessionmanager.session() as session, session.begin():
        one = (await session.execute(text("select 1"))).scalar_one()
    outcome = {
        "refused": refused,
        "enabled_after_refusal": object_content_runtime.enabled,
        "select_one": one,
    }
    await stop_persistence()
    outcome["enabled_after_stop"] = object_content_runtime.enabled
    print(json.dumps(outcome))


asyncio.run(main())
"""


def _run_harness_process(
    test_settings,
    *,
    user_id: UUID,
    space_id: UUID,
    template_file_id: UUID | None,
    fail_first_start: bool = False,
) -> dict[str, Any]:
    payload = {
        "fail_first_start": fail_first_start,
        "user_id": str(user_id),
        "space_id": str(space_id),
        "template_file_id": None if template_file_id is None else str(template_file_id),
        "template_name": None if template_file_id is None else "mall.docx",
        "spec": _spec(template=template_file_id is not None).model_dump(mode="json"),
    }
    completed = subprocess.run(
        [sys.executable, "-c", _HARNESS_PROCESS, str(_BACKEND / "scripts")],
        input=json.dumps(payload),
        capture_output=True,
        text=True,
        timeout=180,
        cwd=_BACKEND,
        env=_flow_worker_environment(settings=test_settings, queue_name="unused"),
    )
    assert completed.returncode == 0, completed.stderr[-4000:]
    return json.loads(completed.stdout.strip().splitlines()[-1])


@dataclass(frozen=True)
class Persisted:
    space_id: UUID
    step_names: list[str | None]
    template_assets: list[tuple[UUID, list[str]]]


async def _persisted(flow_id: UUID) -> Persisted:
    async with sessionmanager.session() as session, session.begin():
        flow = (
            await session.execute(sa.select(Flows).where(Flows.id == flow_id))
        ).scalar_one()
        steps = (
            await session.execute(
                sa.select(FlowSteps.user_description)
                .where(FlowSteps.flow_id == flow_id)
                .order_by(FlowSteps.step_order)
            )
        ).scalars()
        assets = (
            await session.execute(
                sa.select(FlowTemplateAssets.file_id, FlowTemplateAssets.placeholders)
                .where(FlowTemplateAssets.flow_id == flow_id)
                .order_by(FlowTemplateAssets.created_at)
            )
        ).all()
        return Persisted(
            space_id=flow.space_id,
            step_names=list(steps),
            template_assets=[(file_id, list(names)) for file_id, names in assets],
        )


async def test_a_spec_is_applied_in_a_fresh_process_and_the_runtime_is_stopped_after(
    db_container, test_settings
) -> None:
    user_id, space_id, _file_id = await _space_and_template(db_container)

    outcome = _run_harness_process(
        test_settings, user_id=user_id, space_id=space_id, template_file_id=None
    )

    assert outcome["object_content_enabled"] is False
    assert outcome["session_manager_stopped"] is True
    assert len(set(outcome["flow_ids"])) == 2
    for flow_id in outcome["flow_ids"]:
        persisted = await _persisted(UUID(flow_id))
        assert persisted.space_id == space_id
        assert persisted.step_names == ["Läs ärendet", "Skriv beslutet"]
        assert persisted.template_assets == []


async def test_a_template_is_read_through_the_object_content_the_process_started(
    db_container, test_settings
) -> None:
    user_id, space_id, file_id = await _space_and_template(db_container)

    outcome = _run_harness_process(
        test_settings, user_id=user_id, space_id=space_id, template_file_id=file_id
    )

    assert outcome["object_content_enabled"] is False
    assert outcome["session_manager_stopped"] is True
    for flow_id in outcome["flow_ids"]:
        persisted = await _persisted(UUID(flow_id))
        assert persisted.template_assets == [(file_id, ["case_id"])]


async def test_a_failed_start_leaves_nothing_behind_and_the_next_apply_succeeds(
    db_container, test_settings
) -> None:
    user_id, space_id, _file_id = await _space_and_template(db_container)

    outcome = _run_harness_process(
        test_settings,
        user_id=user_id,
        space_id=space_id,
        template_file_id=None,
        fail_first_start=True,
    )

    assert "the store is unreachable" in outcome["first_error"]
    assert outcome["object_content_enabled"] is False
    assert outcome["session_manager_stopped"] is True
    for flow_id in outcome["flow_ids"]:
        assert (await _persisted(UUID(flow_id))).space_id == space_id


async def test_a_second_start_is_refused_and_the_first_owner_stays_usable(
    setup_database, test_settings
) -> None:
    del setup_database
    completed = subprocess.run(
        [sys.executable, "-c", _SECOND_START_PROCESS],
        capture_output=True,
        text=True,
        timeout=120,
        cwd=_BACKEND,
        env=_flow_worker_environment(settings=test_settings, queue_name="unused"),
    )
    assert completed.returncode == 0, completed.stderr[-4000:]

    outcome = json.loads(completed.stdout.strip().splitlines()[-1])
    assert "already running" in outcome["refused"]
    assert outcome["enabled_after_refusal"] is True
    assert outcome["select_one"] == 1
    assert outcome["enabled_after_stop"] is False
