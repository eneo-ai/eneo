import asyncio
import os
import sys
import textwrap
import zipfile
from pathlib import Path

import pytest

from eneo.files import bounded_extraction
from eneo.files.extraction_limits import FileExtractionLimits
from eneo.files.text import ExtractionLimitError, TextExtractor


@pytest.fixture(autouse=True)
def extraction_slots(monkeypatch):
    monkeypatch.setattr(bounded_extraction, "_process_slots", asyncio.Semaphore(2))


@pytest.fixture
def controlled_child(monkeypatch):
    original_spawn = asyncio.create_subprocess_exec

    def replace(code, *, delay_spawn=False):
        processes = []
        started = asyncio.Event()
        release_spawn = asyncio.Event()

        async def spawn(*args, **kwargs):
            process = await original_spawn(
                sys.executable, "-c", code, *args[3:], **kwargs
            )
            processes.append(process)
            started.set()
            if delay_spawn:
                await release_spawn.wait()
            return process

        monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
        return processes, started, release_spawn

    return replace


@pytest.mark.parametrize("suffix", ["txt", "docx", "xlsx", "pptx", "pdf"])
async def test_bounded_extraction_preserves_supported_document_text(
    tmp_path: Path, suffix: str
) -> None:
    path = tmp_path / f"document.{suffix}"
    if suffix == "txt":
        path.write_text("Knowledge in Swedish: räksmörgås", encoding="utf-8")
    elif suffix == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("Knowledge from a document")
        document.save(path)
    elif suffix == "xlsx":
        import pandas as pd

        pd.DataFrame({"Knowledge": ["from a spreadsheet"]}).to_excel(path, index=False)
    elif suffix == "pptx":
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Knowledge from a presentation"
        presentation.save(path)
    else:
        path.write_bytes(
            b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
            b"/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>endobj\n"
            b"4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
            b"5 0 obj<</Length 57>>stream\n"
            b"BT /F1 12 Tf 40 700 Td (Knowledge from a PDF) Tj ET\n"
            b"endstream\nendobj\ntrailer<</Root 1 0 R>>\n%%EOF"
        )

    extractor = TextExtractor()
    expected = extractor.extract(path)
    assert "Knowledge" in expected
    assert await extractor.extract_bounded(path) == expected


@pytest.mark.parametrize(
    "budget", ["output", "archive_bytes", "archive_entries", "pages"]
)
async def test_document_budget_rejects_oversized_input(tmp_path, budget):
    path = tmp_path / "input.txt"
    if budget == "output":
        path.write_text("ä" * 100, encoding="utf-8")
        limits = FileExtractionLimits(max_output_bytes=100)
    elif budget in ("archive_bytes", "archive_entries"):
        path = tmp_path / "input.docx"
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index in range(3):
                archive.writestr(str(index), "x" * 2_000)
        limits = FileExtractionLimits(
            max_archive_bytes=100 if budget == "archive_bytes" else 100_000,
            max_archive_entries=2 if budget == "archive_entries" else 100,
        )
    else:
        from PIL import Image

        path = tmp_path / "input.pdf"
        with Image.new("RGB", (10, 10)) as page:
            page.save(path, save_all=True, append_images=[page])
        limits = FileExtractionLimits(max_pdf_pages=1)

    with pytest.raises(ExtractionLimitError):
        await TextExtractor().extract_bounded(path, limits=limits)


async def test_high_ratio_docx_within_absolute_budgets_is_accepted(tmp_path):
    from docx import Document

    path = tmp_path / "compressible.docx"
    text = "Municipal information with useful searchable content. " * 40_000
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    assert len(text.encode()) > 40 * path.stat().st_size
    assert text.strip() == await TextExtractor().extract_bounded(path)


@pytest.mark.parametrize("stop", ["timeout", "cancel", "cancel_during_spawn"])
async def test_stalled_children_are_reaped_before_returning(
    tmp_path, controlled_child, stop
):
    processes, started, release_spawn = controlled_child(
        "import time; time.sleep(60)", delay_spawn=stop == "cancel_during_spawn"
    )
    for _ in range(3):
        started.clear()
        release_spawn.clear()
        task = asyncio.create_task(
            TextExtractor().extract_bounded(
                tmp_path / "input.txt",
                limits=FileExtractionLimits(
                    timeout_seconds=0.2 if stop == "timeout" else 10
                ),
            )
        )
        await asyncio.wait_for(started.wait(), timeout=5)
        if stop != "timeout":
            task.cancel()
            release_spawn.set()
        expected = ExtractionLimitError if stop == "timeout" else asyncio.CancelledError
        with pytest.raises(expected):
            await asyncio.wait_for(task, timeout=5)
        assert processes[-1].returncode is not None
    assert all(process.returncode is not None for process in processes)


async def test_extraction_child_count_is_bounded_and_waiters_can_cancel(
    tmp_path, controlled_child
):
    processes, started, _ = controlled_child("import time; time.sleep(60)")
    tasks = [
        asyncio.create_task(TextExtractor().extract_bounded(tmp_path / str(index)))
        for index in range(3)
    ]
    try:
        async with asyncio.timeout(5):
            while len(processes) < 2:
                await started.wait()
                started.clear()
        await asyncio.sleep(0.05)
        assert len(processes) == 2
    finally:
        for task in tasks:
            task.cancel()
        results = await asyncio.gather(*tasks, return_exceptions=True)
    assert all(isinstance(result, asyncio.CancelledError) for result in results)
    assert all(process.returncode is not None for process in processes)


async def test_parent_caps_output_even_when_child_does_not_cooperate(
    tmp_path, controlled_child
):
    processes, _, _ = controlled_child(
        "import sys; sys.stdout.buffer.write(b'x' * 10_000_000)"
    )
    with pytest.raises(ExtractionLimitError):
        await asyncio.wait_for(
            TextExtractor().extract_bounded(
                tmp_path / "input.txt",
                limits=FileExtractionLimits(max_output_bytes=1_000),
            ),
            timeout=5,
        )
    assert processes[0].returncode is not None


@pytest.mark.skipif(
    sys.platform != "linux", reason="Linux enforces address-space limits"
)
@pytest.mark.parametrize("resource", ["cpu", "memory"])
async def test_linux_child_resource_limits_leave_other_extractions_usable(
    tmp_path, controlled_child, monkeypatch, resource
):
    original_spawn = asyncio.create_subprocess_exec
    code = (
        "import sys; "
        "from eneo.files.extraction_worker import configure_process_limits, ExtractionExit; "
        "configure_process_limits(*map(int, sys.argv[1:4])); "
    )
    code += (
        "\nwhile True: pass"
        if resource == "cpu"
        else "\ntry:\n bytearray(256_000_000)\nexcept MemoryError:\n sys.exit(ExtractionExit.LIMIT)"
    )
    processes, _, _ = controlled_child(code)
    with pytest.raises(ExtractionLimitError):
        await TextExtractor().extract_bounded(
            tmp_path / "input.txt",
            limits=FileExtractionLimits(
                cpu_seconds=1, memory_bytes=67_108_864, timeout_seconds=5
            ),
        )
    assert processes[0].returncode is not None
    monkeypatch.setattr(asyncio, "create_subprocess_exec", original_spawn)
    path = tmp_path / "normal.txt"
    path.write_text("Other extraction still works")
    assert await TextExtractor().extract_bounded(path) == path.read_text()


@pytest.mark.skipif(sys.platform != "linux", reason="Linux parent-death signal")
async def test_linux_extractor_stops_when_its_worker_is_killed(tmp_path):
    fifo = tmp_path / "stalled.txt"
    os.mkfifo(fifo)
    parent_code = textwrap.dedent("""
        import asyncio
        from pathlib import Path
        import sys
        from eneo.files.text import TextExtractor
        original_spawn = asyncio.create_subprocess_exec
        async def record_child(*args, **kwargs):
            child = await original_spawn(*args, **kwargs)
            print(child.pid, flush=True)
            return child
        asyncio.create_subprocess_exec = record_child
        asyncio.run(TextExtractor().extract_bounded(Path(sys.argv[1]), "text/plain"))
    """)
    parent = await asyncio.create_subprocess_exec(
        sys.executable,
        "-c",
        parent_code,
        str(fifo),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        assert parent.stdout is not None
        child_pid = int(await asyncio.wait_for(parent.stdout.readline(), timeout=5))
        parent.kill()
        await parent.wait()
        async with asyncio.timeout(5):
            while True:
                try:
                    state = Path(f"/proc/{child_pid}/stat").read_text().split()[2]
                except FileNotFoundError:
                    break
                # Reaping after a worker SIGKILL belongs to the container init.
                if state == "Z":
                    break
                await asyncio.sleep(0.01)
    finally:
        if parent.returncode is None:
            parent.kill()
            await parent.wait()
