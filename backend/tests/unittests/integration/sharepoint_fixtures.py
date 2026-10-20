"""Builders for the file formats SharePoint sync downloads.

A plain module: import what you need, nothing is auto-loaded.
"""

import io
from pathlib import Path

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
OCTET = "application/octet-stream"


def docx_bytes(*paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def header_only_xlsx_bytes() -> bytes:
    """A workbook whose only text is its first row, which parsers read as the header."""
    return xlsx_bytes([["Name", "Value"]])


def pptx_bytes(*titles: str) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    for title in titles:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def pdf_bytes(page_texts: list[str]) -> bytes:
    """A minimal PDF with one text line per page; an empty string gives a page without text."""
    page_ids = [3 + 2 * index for index in range(len(page_texts))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    font_id = page_ids[-1] + 2
    objects = [
        "1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj",
        f"2 0 obj<</Type/Pages/Kids[{kids}]/Count {len(page_texts)}>>endobj",
        f"{font_id} 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj",
    ]
    for page_id, text in zip(page_ids, page_texts):
        stream = f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET"
        objects.append(
            f"{page_id} 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
            f"/Resources<</Font<</F1 {font_id} 0 R>>>>"
            f"/Contents {page_id + 1} 0 R>>endobj"
        )
        objects.append(
            f"{page_id + 1} 0 obj<</Length {len(stream)}>>stream\n"
            f"{stream}\nendstream\nendobj"
        )
    body = "\n".join(objects)
    return f"%PDF-1.4\n{body}\ntrailer<</Root 1 0 R>>\n%%EOF".encode()


def text_pdf(path: Path, page_texts: list[str]) -> Path:
    path.write_bytes(pdf_bytes(page_texts))
    return path
