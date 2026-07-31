from __future__ import annotations

import hashlib
import io
from datetime import datetime, timezone
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from docx import Document
from fastapi import UploadFile

from eneo.authentication.principal_types import PrincipalType
from eneo.files.file_models import (
    File,
    FileContentVariant,
    FileInfo,
    FileMetadata,
    FileType,
)
from eneo.files.file_protocol import PendingFileContent, PreparedFileUpload
from eneo.flows.domain.flow import Flow, FlowTemplateAsset
from eneo.flows.domain.flow_invariant_exceptions import FlowPersistedIdMissingError
from eneo.flows.flow_template_asset_service import (
    AttachedTemplateFileUnavailableError,
    FlowTemplateAssetService,
)
from eneo.main.exceptions import (
    BadRequestException,
    ConflictException,
    FileTooLargeException,
    NotFoundException,
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class _ReadGuard(io.BytesIO):
    def read(self, *args, **kwargs):
        raise AssertionError("upload body was read before size validation")


def _build_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("{{Body}}")
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


def _upload(filename: str, content: bytes) -> UploadFile:
    return UploadFile(
        file=io.BytesIO(content),
        filename=filename,
        headers={"content-type": DOCX_MIME},
    )


async def _chunks(payload: bytes):
    yield payload


def _prepared_template(
    payload: bytes,
    *,
    filename: str = "template.docx",
) -> PreparedFileUpload:
    return PreparedFileUpload(
        name=filename,
        file_type=FileType.DOCUMENT,
        display_media_type=DOCX_MIME,
        contents=(
            PendingFileContent(
                variant=FileContentVariant.ORIGINAL,
                chunks=_chunks(payload),
                declared_media_type=DOCX_MIME,
                verified_media_type=DOCX_MIME,
            ),
        ),
    )


def _prepared_context(prepared: PreparedFileUpload) -> MagicMock:
    context = MagicMock()
    context.__aenter__ = AsyncMock(return_value=prepared)
    context.__aexit__ = AsyncMock(return_value=False)
    return context


def _saved_file(user, payload: bytes, *, filename: str = "template.docx") -> FileInfo:
    now = datetime.now(timezone.utc)
    return FileInfo(
        id=uuid4(),
        created_at=now,
        updated_at=now,
        name=filename,
        checksum=hashlib.sha256(payload).hexdigest(),
        size=len(payload),
        mimetype=DOCX_MIME,
        file_type=FileType.DOCUMENT,
        owner_type=PrincipalType.USER,
        owner_user_id=user.id,
        owner_service_id=None,
        tenant_id=user.tenant_id,
    )


def _flow_for_user(user) -> Flow:
    now = datetime.now(timezone.utc)
    return Flow(
        id=uuid4(),
        tenant_id=user.tenant_id,
        space_id=uuid4(),
        name="Template flow",
        description=None,
        created_by_user_id=user.id,
        owner_user_id=user.id,
        published_version=None,
        metadata_json=None,
        data_retention_days=None,
        created_at=now,
        updated_at=now,
        steps=[],
    )


def _asset_for_flow(flow: Flow, user) -> FlowTemplateAsset:
    now = datetime.now(timezone.utc)
    assert flow.id is not None
    return FlowTemplateAsset(
        id=uuid4(),
        flow_id=flow.id,
        space_id=flow.space_id,
        tenant_id=user.tenant_id,
        file_id=uuid4(),
        name="template.docx",
        checksum="checksum",
        mimetype=DOCX_MIME,
        placeholders=["Body"],
        created_by_user_id=user.id,
        updated_by_user_id=user.id,
        status="ready",
        created_at=now,
        updated_at=now,
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("operation", ["list", "upload", "get", "delete"])
async def test_template_asset_operations_require_persisted_parent_flow_id(
    user,
    operation: str,
) -> None:
    flow = _flow_for_user(user).model_copy(update={"id": None})
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    file_repo = AsyncMock()
    file_service = AsyncMock()
    template_asset_repo = AsyncMock()
    flow_version_repo = AsyncMock()
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=file_repo,
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=flow_version_repo,
    )

    with pytest.raises(FlowPersistedIdMissingError):
        if operation == "list":
            await service.list_assets(flow_id=uuid4())
        elif operation == "upload":
            await service.upload_asset(
                flow_id=uuid4(),
                upload_file=_upload("template.docx", _build_template_bytes()),
            )
        elif operation == "delete":
            await service.delete_asset(flow_id=uuid4(), asset_id=uuid4())
        else:
            await service.get_asset_with_file(flow_id=uuid4(), asset_id=uuid4())

    template_asset_repo.list_for_flow.assert_not_awaited()
    template_asset_repo.create.assert_not_awaited()
    template_asset_repo.get.assert_not_awaited()
    template_asset_repo.soft_delete.assert_not_awaited()
    flow_version_repo.has_template_asset_reference.assert_not_awaited()
    file_service.prepare_document_upload.assert_not_called()


@pytest.mark.asyncio
async def test_get_asset_with_file_hydrates_relationship_content(user) -> None:
    flow = _flow_for_user(user)
    asset = _asset_for_flow(flow, user)
    now = datetime.now(timezone.utc)
    metadata = FileMetadata(
        id=asset.file_id,
        created_at=now,
        updated_at=now,
        name=asset.name,
        file_type=FileType.DOCUMENT,
        mimetype=asset.mimetype,
        owner_type=PrincipalType.USER,
        owner_user_id=user.id,
        owner_service_id=None,
        tenant_id=user.tenant_id,
    )
    template_bytes = _build_template_bytes()
    hydrated_file = File(
        id=metadata.id,
        created_at=metadata.created_at,
        updated_at=metadata.updated_at,
        name=metadata.name,
        checksum=hashlib.sha256(template_bytes).hexdigest(),
        size=len(template_bytes),
        mimetype=metadata.mimetype,
        file_type=metadata.file_type,
        blob=template_bytes,
        owner_type=metadata.owner_type,
        owner_user_id=metadata.owner_user_id,
        owner_service_id=metadata.owner_service_id,
        tenant_id=metadata.tenant_id,
    )
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()
    template_asset_repo.get.return_value = asset
    file_repo = AsyncMock()
    file_repo.get_by_id.return_value = metadata
    file_content_loader = AsyncMock()
    file_content_loader.load.return_value = {metadata.id: hydrated_file}
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=file_repo,
        file_content_loader=file_content_loader,
        file_service=AsyncMock(),
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )

    result = await service.get_asset_with_file(flow_id=flow.id, asset_id=asset.id)

    assert result == (asset, hydrated_file)
    file_repo.get_by_id.assert_awaited_once_with(
        file_id=asset.file_id,
        tenant_id=user.tenant_id,
    )
    file_content_loader.load.assert_awaited_once_with(
        [metadata],
        include_text_original_bytes=True,
    )


@pytest.mark.asyncio
async def test_upload_asset_persists_docx_template_bytes_and_body_placeholder(
    user,
) -> None:
    template_bytes = _build_template_bytes()
    flow = _flow_for_user(user)
    file_repo = AsyncMock()
    now = datetime.now(timezone.utc)
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()

    async def create_asset(**kwargs) -> FlowTemplateAsset:
        return FlowTemplateAsset(
            id=uuid4(),
            created_at=now,
            updated_at=now,
            **kwargs,
        )

    template_asset_repo.create.side_effect = create_asset
    prepared = _prepared_template(template_bytes)
    saved_file = _saved_file(user, template_bytes)
    file_service = MagicMock()
    file_service.prepare_document_upload.return_value = _prepared_context(prepared)
    file_service.save_prepared_file = AsyncMock(return_value=saved_file)
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=file_repo,
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )

    upload = _upload("template.docx", template_bytes)
    asset = await service.upload_asset(
        flow_id=flow.id,
        upload_file=upload,
    )

    file_service.prepare_document_upload.assert_called_once_with(upload)
    file_service.save_prepared_file.assert_awaited_once_with(prepared)
    template_asset_repo.create.assert_awaited_once()
    assert template_asset_repo.create.await_args.kwargs["file_id"] == saved_file.id
    assert template_asset_repo.create.await_args.kwargs["checksum"] == (
        saved_file.checksum
    )
    assert template_asset_repo.create.await_args.kwargs["placeholders"] == ["Body"]
    assert asset.placeholders == ["Body"]


@pytest.mark.asyncio
async def test_upload_asset_rejects_invalid_docx_before_file_persistence(
    user,
) -> None:
    flow = _flow_for_user(user)
    file_repo = AsyncMock()
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()
    prepared = _prepared_template(b"not-a-docx")
    file_service = MagicMock()
    file_service.prepare_document_upload.return_value = _prepared_context(prepared)
    file_service.save_prepared_file = AsyncMock()
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=file_repo,
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )

    with pytest.raises(BadRequestException) as exc_info:
        await service.upload_asset(
            flow_id=flow.id,
            upload_file=_upload("template.docx", b"not-a-docx"),
        )

    assert exc_info.value.code == "flow_template_invalid_archive"
    file_service.save_prepared_file.assert_not_awaited()
    template_asset_repo.create.assert_not_awaited()


@pytest.mark.asyncio
async def test_create_from_existing_attached_file_reuses_authorized_file(user) -> None:
    template_bytes = _build_template_bytes()
    flow = _flow_for_user(user)
    now = datetime.now(timezone.utc)
    attached_file = File(
        id=uuid4(),
        name="template.docx",
        checksum=hashlib.sha256(template_bytes).hexdigest(),
        size=len(template_bytes),
        mimetype=DOCX_MIME,
        file_type=FileType.DOCUMENT,
        blob=template_bytes,
        owner_type=PrincipalType.USER,
        owner_user_id=user.id,
        tenant_id=user.tenant_id,
        created_at=now,
        updated_at=now,
    )
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    file_service = AsyncMock()
    file_service.get_owned_file_for_key_share.return_value = attached_file
    template_asset_repo = AsyncMock()
    template_asset_repo.find_ready_for_file.return_value = None

    async def create_asset(**kwargs) -> FlowTemplateAsset:
        return FlowTemplateAsset(
            id=uuid4(),
            created_at=now,
            updated_at=now,
            **kwargs,
        )

    template_asset_repo.create.side_effect = create_asset
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )

    asset = await service.create_from_existing_attached_file(
        flow_id=flow.id,
        file_id=attached_file.id,
    )

    file_service.get_owned_file_for_key_share.assert_awaited_once_with(
        attached_file.id,
        include_text_original_bytes=True,
    )
    file_service.save_file_content.assert_not_awaited()
    template_asset_repo.create.assert_awaited_once_with(
        flow_id=flow.id,
        space_id=flow.space_id,
        tenant_id=user.tenant_id,
        file_id=attached_file.id,
        name=attached_file.name,
        checksum=attached_file.checksum,
        mimetype=attached_file.mimetype,
        placeholders=["Body"],
        created_by_user_id=user.id,
        updated_by_user_id=user.id,
        status="ready",
    )
    assert asset.file_id == attached_file.id
    assert asset.placeholders == ["Body"]


@pytest.mark.asyncio
async def test_create_from_existing_attached_file_rejects_missing_content(user) -> None:
    flow = _flow_for_user(user)
    attached_file = MagicMock()
    attached_file.id = uuid4()
    attached_file.blob = None
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    file_service = AsyncMock()
    file_service.get_owned_file_for_key_share.return_value = attached_file
    template_asset_repo = AsyncMock()
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )

    with pytest.raises(BadRequestException) as exc_info:
        await service.create_from_existing_attached_file(
            flow_id=flow.id,
            file_id=attached_file.id,
        )

    assert exc_info.value.code == "flow_template_missing_content"
    template_asset_repo.create.assert_not_awaited()


@pytest.mark.asyncio
async def test_create_from_existing_attached_file_reuses_active_same_file_asset(
    user,
) -> None:
    template_bytes = _build_template_bytes()
    flow = _flow_for_user(user)
    now = datetime.now(timezone.utc)
    attached_file = File(
        id=uuid4(),
        name="template.docx",
        checksum=hashlib.sha256(template_bytes).hexdigest(),
        size=len(template_bytes),
        mimetype=DOCX_MIME,
        file_type=FileType.DOCUMENT,
        blob=template_bytes,
        owner_type=PrincipalType.USER,
        owner_user_id=user.id,
        tenant_id=user.tenant_id,
        created_at=now,
        updated_at=now,
    )
    reusable = _asset_for_flow(flow, user).model_copy(
        update={
            "file_id": attached_file.id,
            "checksum": attached_file.checksum,
        }
    )
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    file_service = AsyncMock()
    file_service.get_owned_file_for_key_share.return_value = attached_file
    template_asset_repo = AsyncMock()
    template_asset_repo.find_ready_for_file.return_value = reusable
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )

    result = await service.create_from_existing_attached_file(
        flow_id=flow.require_persisted_id(),
        file_id=attached_file.id,
    )

    assert result == reusable
    template_asset_repo.find_ready_for_file.assert_awaited_once_with(
        flow_id=flow.require_persisted_id(),
        tenant_id=user.tenant_id,
        file_id=attached_file.id,
        checksum=attached_file.checksum,
    )
    template_asset_repo.create.assert_not_awaited()


@pytest.mark.asyncio
async def test_create_from_existing_attached_file_reports_deleted_file(user) -> None:
    flow = _flow_for_user(user)
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    file_service = AsyncMock()
    file_service.get_owned_file_for_key_share.side_effect = NotFoundException()
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=AsyncMock(),
        flow_version_repo=AsyncMock(),
    )

    with pytest.raises(AttachedTemplateFileUnavailableError):
        await service.create_from_existing_attached_file(
            flow_id=flow.require_persisted_id(),
            file_id=uuid4(),
        )


@pytest.mark.asyncio
async def test_upload_asset_rejects_oversized_docx_before_reading_body(
    user,
) -> None:
    flow = _flow_for_user(user)
    file_repo = AsyncMock()
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()
    file_service = MagicMock()
    rejected_context = MagicMock()
    rejected_context.__aenter__ = AsyncMock(
        side_effect=FileTooLargeException(
            file_size=10**12,
            max_size=10,
            limit_name="session_file",
        )
    )
    rejected_context.__aexit__ = AsyncMock(return_value=False)
    file_service.prepare_document_upload.return_value = rejected_context
    file_service.save_prepared_file = AsyncMock()
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=file_repo,
        file_content_loader=AsyncMock(),
        file_service=file_service,
        template_asset_repo=template_asset_repo,
        flow_version_repo=AsyncMock(),
    )
    upload = UploadFile(
        file=_ReadGuard(b"body-must-not-be-read"),
        filename="oversized.docx",
        headers={"content-type": DOCX_MIME},
    )

    with pytest.raises(FileTooLargeException):
        await service.upload_asset(flow_id=flow.id, upload_file=upload)

    file_service.prepare_document_upload.assert_called_once_with(upload)
    file_service.save_prepared_file.assert_not_awaited()
    template_asset_repo.create.assert_not_awaited()


@pytest.mark.asyncio
async def test_delete_asset_soft_deletes_unpinned_template_asset(user) -> None:
    flow = _flow_for_user(user)
    asset = _asset_for_flow(flow, user)
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()
    template_asset_repo.get.return_value = asset
    flow_version_repo = AsyncMock()
    flow_version_repo.has_template_asset_reference.return_value = False
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=AsyncMock(),
        template_asset_repo=template_asset_repo,
        flow_version_repo=flow_version_repo,
    )

    deleted = await service.delete_asset(flow_id=flow.id, asset_id=asset.id)

    assert deleted == asset
    flow_version_repo.has_template_asset_reference.assert_awaited_once_with(
        flow_id=flow.id,
        tenant_id=user.tenant_id,
        template_asset_id=asset.id,
        template_file_id=asset.file_id,
    )
    template_asset_repo.soft_delete.assert_awaited_once_with(
        asset_id=asset.id,
        tenant_id=user.tenant_id,
        updated_by_user_id=user.id,
    )


@pytest.mark.asyncio
async def test_delete_asset_rejects_published_definition_pin(user) -> None:
    flow = _flow_for_user(user)
    asset = _asset_for_flow(flow, user)
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()
    template_asset_repo.get.return_value = asset
    flow_version_repo = AsyncMock()
    flow_version_repo.has_template_asset_reference.return_value = True
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=AsyncMock(),
        template_asset_repo=template_asset_repo,
        flow_version_repo=flow_version_repo,
    )

    with pytest.raises(ConflictException) as exc_info:
        await service.delete_asset(flow_id=flow.id, asset_id=asset.id)

    assert exc_info.value.code == "flow_template_in_use"
    assert exc_info.value.context == {
        "flow_id": str(flow.id),
        "template_asset_id": str(asset.id),
        "template_file_id": str(asset.file_id),
    }
    template_asset_repo.soft_delete.assert_not_awaited()


@pytest.mark.asyncio
async def test_delete_asset_hides_wrong_flow_asset(user) -> None:
    flow = _flow_for_user(user)
    other_flow = _flow_for_user(user)
    asset = _asset_for_flow(other_flow, user)
    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    template_asset_repo = AsyncMock()
    template_asset_repo.get.return_value = asset
    flow_version_repo = AsyncMock()
    service = FlowTemplateAssetService(
        user=user,
        flow_repo=flow_repo,
        file_repo=AsyncMock(),
        file_content_loader=AsyncMock(),
        file_service=AsyncMock(),
        template_asset_repo=template_asset_repo,
        flow_version_repo=flow_version_repo,
    )

    with pytest.raises(NotFoundException):
        await service.delete_asset(flow_id=flow.id, asset_id=asset.id)

    flow_version_repo.has_template_asset_reference.assert_not_awaited()
    template_asset_repo.soft_delete.assert_not_awaited()
