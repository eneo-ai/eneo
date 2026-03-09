from __future__ import annotations

import io
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from docx import Document

from intric.flows.flow import Flow, FlowStep
from intric.flows.flow_file_upload_service import FlowFileUploadService
from intric.flows.flow_input_limits import FlowInputLimits
from intric.flows.flow_service import FlowService
from intric.main.exceptions import BadRequestException, NotFoundException


def _build_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("{{Body}}")
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


def _step(step_order: int = 1) -> FlowStep:
    return FlowStep(
        id=uuid4(),
        assistant_id=uuid4(),
        step_order=step_order,
        user_description=f"Step {step_order}",
        input_source="flow_input" if step_order == 1 else "previous_step",
        input_type="text",
        output_mode="pass_through",
        output_type="json",
        mcp_policy="inherit",
    )


def _flow(*steps: FlowStep) -> Flow:
    now = datetime.now(timezone.utc)
    return Flow(
        id=uuid4(),
        tenant_id=uuid4(),
        space_id=uuid4(),
        name="Flow",
        description=None,
        created_by_user_id=uuid4(),
        owner_user_id=uuid4(),
        published_version=None,
        metadata_json=None,
        data_retention_days=None,
        created_at=now,
        updated_at=now,
        steps=list(steps),
    )


@pytest.mark.asyncio
async def test_publish_flow_resolves_legacy_template_file_id_through_flow_asset(user) -> None:
    template_file_id = uuid4()
    template_asset_id = uuid4()
    template_step = _step().model_copy(
        update={
            "output_mode": "template_fill",
            "output_type": "docx",
            "output_config": {
                "template_file_id": str(template_file_id),
                "bindings": {"Body": "{{flow_input.title}}"},
            },
        }
    )
    flow = _flow(template_step).model_copy(
        update={
            "id": uuid4(),
            "tenant_id": user.tenant_id,
            "created_by_user_id": user.id,
            "owner_user_id": user.id,
        }
    )

    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    flow_repo.update.return_value = flow.model_copy(update={"published_version": 1})
    version_repo = AsyncMock()
    version_repo.get_latest.return_value = None
    assistant_service = AsyncMock()
    file_repo = AsyncMock()
    template_asset_repo = AsyncMock()
    template_asset_repo.get_by_flow_file.return_value = SimpleNamespace(
        id=template_asset_id,
        flow_id=flow.id,
        file_id=template_file_id,
        name="mall.docx",
        checksum="checksum-1",
    )
    template_asset_repo.get.return_value = SimpleNamespace(
        id=template_asset_id,
        flow_id=flow.id,
        file_id=template_file_id,
        name="mall.docx",
        checksum="checksum-1",
    )
    file_repo.get_by_id.return_value = SimpleNamespace(
        id=template_file_id,
        name="mall.docx",
        checksum="checksum-1",
        tenant_id=user.tenant_id,
        blob=b"docx-bytes",
    )

    service = FlowService(
        user=user,
        flow_repo=flow_repo,
        flow_version_repo=version_repo,
        assistant_service=assistant_service,
        file_repo=file_repo,
        template_asset_repo=template_asset_repo,
    )
    service._validate_assistant_scope_for_steps = AsyncMock()  # type: ignore[method-assign]
    service._inspect_docx_template = MagicMock(  # type: ignore[method-assign]
        return_value=[{"name": "Body", "location": "body"}]
    )

    await service.publish_flow(flow_id=flow.id)

    definition_json = version_repo.create.await_args.kwargs["definition_json"]
    published_step = definition_json["steps"][0]
    assert published_step["output_config"]["template_asset_id"] == str(template_asset_id)
    assert published_step["output_config"]["template_file_id"] == str(template_file_id)
    template_asset_repo.get_by_flow_file.assert_awaited_once_with(
        flow_id=flow.id,
        file_id=template_file_id,
        tenant_id=user.tenant_id,
    )


@pytest.mark.asyncio
async def test_publish_flow_promotes_legacy_template_file_id_into_flow_asset_when_missing(user) -> None:
    template_file_id = uuid4()
    created_asset_id = uuid4()
    template_step = _step().model_copy(
        update={
            "output_mode": "template_fill",
            "output_type": "docx",
            "output_config": {
                "template_file_id": str(template_file_id),
                "bindings": {"Body": "{{flow_input.title}}"},
            },
        }
    )
    flow = _flow(template_step).model_copy(
        update={
            "id": uuid4(),
            "tenant_id": user.tenant_id,
            "space_id": uuid4(),
            "created_by_user_id": user.id,
            "owner_user_id": user.id,
        }
    )

    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    flow_repo.update.return_value = flow.model_copy(update={"published_version": 1})
    version_repo = AsyncMock()
    version_repo.get_latest.return_value = None
    assistant_service = AsyncMock()
    file_repo = AsyncMock()
    template_asset_repo = AsyncMock()
    template_asset_repo.get_by_flow_file.side_effect = NotFoundException()
    template_asset_repo.create.return_value = SimpleNamespace(
        id=created_asset_id,
        flow_id=flow.id,
        file_id=template_file_id,
        name="legacy-template.docx",
        checksum="legacy-checksum",
    )
    template_asset_repo.get.return_value = SimpleNamespace(
        id=created_asset_id,
        flow_id=flow.id,
        file_id=template_file_id,
        name="legacy-template.docx",
        checksum="legacy-checksum",
    )
    file_repo.get_by_id.return_value = SimpleNamespace(
        id=template_file_id,
        name="legacy-template.docx",
        checksum="legacy-checksum",
        mimetype="application/octet-stream",
        tenant_id=user.tenant_id,
        blob=_build_template_bytes(),
    )

    service = FlowService(
        user=user,
        flow_repo=flow_repo,
        flow_version_repo=version_repo,
        assistant_service=assistant_service,
        file_repo=file_repo,
        template_asset_repo=template_asset_repo,
    )
    service._validate_assistant_scope_for_steps = AsyncMock()  # type: ignore[method-assign]

    await service.publish_flow(flow_id=flow.id)

    template_asset_repo.create.assert_awaited_once_with(
        flow_id=flow.id,
        space_id=flow.space_id,
        tenant_id=user.tenant_id,
        file_id=template_file_id,
        name="legacy-template.docx",
        checksum="legacy-checksum",
        mimetype="application/octet-stream",
        placeholders=["Body"],
        created_by_user_id=user.id,
        updated_by_user_id=user.id,
        status="ready",
    )
    definition_json = version_repo.create.await_args.kwargs["definition_json"]
    published_step = definition_json["steps"][0]
    assert published_step["output_config"]["template_asset_id"] == str(created_asset_id)
    assert published_step["output_config"]["template_file_id"] == str(template_file_id)


@pytest.mark.asyncio
async def test_publish_flow_reports_missing_legacy_template_file_as_bad_request(user) -> None:
    template_file_id = uuid4()
    template_step = _step().model_copy(
        update={
            "output_mode": "template_fill",
            "output_type": "docx",
            "output_config": {
                "template_file_id": str(template_file_id),
                "bindings": {"Body": "{{flow_input.title}}"},
            },
        }
    )
    flow = _flow(template_step).model_copy(
        update={
            "id": uuid4(),
            "tenant_id": user.tenant_id,
            "space_id": uuid4(),
            "created_by_user_id": user.id,
            "owner_user_id": user.id,
        }
    )

    flow_repo = AsyncMock()
    flow_repo.get.return_value = flow
    version_repo = AsyncMock()
    assistant_service = AsyncMock()
    file_repo = AsyncMock()
    file_repo.get_by_id.side_effect = NotFoundException()
    template_asset_repo = AsyncMock()
    template_asset_repo.get_by_flow_file.side_effect = NotFoundException()

    service = FlowService(
        user=user,
        flow_repo=flow_repo,
        flow_version_repo=version_repo,
        assistant_service=assistant_service,
        file_repo=file_repo,
        template_asset_repo=template_asset_repo,
    )
    service._validate_assistant_scope_for_steps = AsyncMock()  # type: ignore[method-assign]

    with pytest.raises(BadRequestException, match="selected DOCX template is no longer available"):
        await service.publish_flow(flow_id=flow.id)


@pytest.mark.asyncio
async def test_run_contract_marks_legacy_template_file_selection_as_ready_when_asset_exists() -> None:
    flow_service = AsyncMock()
    file_service = AsyncMock()
    settings_service = AsyncMock()
    flow_version_repo = AsyncMock()
    template_asset_repo = AsyncMock()

    template_file_id = uuid4()
    template_asset_id = uuid4()
    template_step = _step(step_order=1).model_copy(
        update={
            "output_mode": "template_fill",
            "output_type": "docx",
            "output_config": {
                "template_file_id": str(template_file_id),
                "template_checksum": "published-checksum",
                "template_name": "Legacy selected template",
                "bindings": {"Body": "{{flow_input.title}}"},
            },
        }
    )
    flow = _flow(template_step).model_copy(update={"published_version": 3})
    flow_service.get_flow.return_value = flow
    settings_service.get_flow_input_limits_resolved.return_value = FlowInputLimits(
        file_max_size_bytes=12_000_000,
        audio_max_size_bytes=25_000_000,
        max_files_per_run=5,
    )
    flow_version_repo.get.return_value = SimpleNamespace(
        definition_json={
            "steps": [
                {
                    "step_id": str(template_step.id),
                    "step_order": 1,
                    "assistant_id": str(template_step.assistant_id),
                    "input_source": "flow_input",
                    "input_type": "text",
                    "output_mode": "template_fill",
                    "output_type": "docx",
                    "output_config": template_step.output_config,
                    "mcp_policy": "inherit",
                }
            ]
        }
    )
    template_asset_repo.get_by_flow_file.return_value = SimpleNamespace(
        id=template_asset_id,
        file_id=template_file_id,
        name="Legacy selected template",
        checksum="published-checksum",
    )

    service = FlowFileUploadService(
        flow_service=flow_service,
        file_service=file_service,
        settings_service=settings_service,
        flow_version_repo=flow_version_repo,
        template_asset_repo=template_asset_repo,
    )

    contract = await service.get_run_contract(flow_id=flow.id)

    assert contract["template_readiness"][0]["status"] == "ready"
    assert contract["template_readiness"][0]["template_asset_id"] == template_asset_id
    template_asset_repo.get_by_flow_file.assert_awaited_once_with(
        flow_id=flow.id,
        file_id=template_file_id,
        tenant_id=flow.tenant_id,
    )
