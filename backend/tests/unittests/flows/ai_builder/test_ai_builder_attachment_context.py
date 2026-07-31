from __future__ import annotations

import io
import json
from typing import get_args
from uuid import UUID, uuid4

import pytest
from docx import Document

from eneo.authentication.principal_types import PrincipalType
from eneo.files.file_models import File, FileType
from eneo.flows.ai_builder.ai_builder_attachment_context import (
    _FILE_ROLE_PRIORITY,
    AIBuilderAttachmentContext,
    AIBuilderAttachmentContextPolicy,
    apply_attachment_structural_evidence_to_planning_state,
    attachment_file_roles,
    build_ai_builder_attachment_context,
    build_ai_builder_attachment_context_for_model,
    render_ai_builder_attachment_evidence,
    render_ai_builder_evidence_value,
)
from eneo.flows.ai_builder.ai_builder_discovery_runtime import (
    build_slot_classification_input,
)
from eneo.flows.ai_builder.ai_builder_error_contract import (
    AIBuilderBadRequestException,
    AIBuilderErrorCode,
)
from eneo.flows.ai_builder.ai_builder_output_schema_evidence import (
    OUTPUT_SCHEMA_MAX_DEPTH,
    OUTPUT_SCHEMA_MAX_JSON_BYTES,
)
from eneo.flows.ai_builder.planning_state import (
    TEMPLATE_PLACEHOLDER_EVIDENCE_PREFIX,
    FileRole,
    PlanningState,
)


def _make_file(
    *,
    name: str,
    text: str | None,
    mimetype: str = "text/plain",
    file_type: FileType = FileType.TEXT,
    transcription: str | None = None,
    file_id: UUID | None = None,
    blob: bytes | None = None,
) -> File:
    readable = text or transcription or ""
    return File(
        id=file_id or uuid4(),
        name=name,
        checksum="checksum",
        size=len(readable.encode("utf-8")) or 1,
        mimetype=mimetype,
        file_type=file_type,
        text=text,
        blob=blob if blob is not None else (b"x" if text is None else None),
        transcription=transcription,
        owner_type=PrincipalType.USER,
        owner_user_id=uuid4(),
        tenant_id=uuid4(),
    )


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


def _build_with_text_budget(
    files: list[File], max_text_chars: int
) -> AIBuilderAttachmentContext | None:
    full_context = build_ai_builder_attachment_context(files)
    assert full_context is not None
    assert full_context.context is not None
    framing_chars = len(full_context.context) - full_context.total_chars
    return build_ai_builder_attachment_context(
        files,
        fits_context=lambda context: context is None
        or len(context) <= framing_chars + max_text_chars,
    )


def test_build_ai_builder_attachment_context_fits_text_fairly() -> None:
    files = [
        _make_file(name="one.txt", text="A" * 120),
        _make_file(name="two.txt", text="B" * 120),
    ]

    result = _build_with_text_budget(files, 70)

    assert result is not None
    assert "one.txt" in result.context
    assert "two.txt" in result.context
    assert len(result.included_file_ids) == 2
    assert result.truncated is True
    assert result.total_chars <= 70


def test_attachment_text_admission_scales_with_selected_model_context() -> None:
    files = [
        _make_file(name="one.txt", text="first evidence " * 2_000),
        _make_file(name="two.txt", text="second evidence " * 2_000),
    ]
    policy = AIBuilderAttachmentContextPolicy()

    small = build_ai_builder_attachment_context_for_model(
        files,
        policy=policy,
        model_name="gpt-4o-mini",
        max_input_tokens=2_000,
        max_output_tokens=500,
        safety_buffer_tokens=250,
        minimum_conversation_tokens=500,
    )
    large = build_ai_builder_attachment_context_for_model(
        files,
        policy=policy,
        model_name="gpt-4o-mini",
        max_input_tokens=8_000,
        max_output_tokens=500,
        safety_buffer_tokens=250,
        minimum_conversation_tokens=500,
    )

    assert small is not None
    assert large is not None
    assert small.truncated is True
    assert large.total_chars > small.total_chars
    assert (
        abs(len(large.evidence[0].excerpt or "") - len(large.evidence[1].excerpt or ""))
        <= 1
    )


def test_file_role_priority_covers_all_declared_file_roles() -> None:
    assert set(_FILE_ROLE_PRIORITY) == set(get_args(FileRole))


def test_build_ai_builder_attachment_context_includes_typed_file_evidence() -> None:
    files = [
        _make_file(
            name="beslutsmall.docx",
            text="Beslutspunkt: ...",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_type=FileType.DOCUMENT,
        ),
        _make_file(
            name="meeting.m4a",
            text=None,
            mimetype="audio/mp4",
            file_type=FileType.AUDIO,
            transcription="Vi beslutade att följa upp avtalet.",
        ),
    ]

    result = build_ai_builder_attachment_context(files)

    assert result is not None
    assert result.evidence[0].file_id == files[0].id
    assert result.evidence[0].filename == "beslutsmall.docx"
    assert result.evidence[0].file_type == FileType.DOCUMENT
    assert result.evidence[0].inferred_role == "context_only"
    assert result.evidence[0].role_confidence == "low"
    assert result.evidence[0].coverage == "fully_seen"
    assert result.evidence[1].file_id == files[1].id
    assert result.evidence[1].filename == "meeting.m4a"
    assert result.evidence[1].file_type == FileType.AUDIO
    assert result.evidence[1].inferred_role == "runtime_input_sample"
    assert result.evidence[1].role_confidence == "high"
    assert result.evidence[1].coverage == "fully_seen"


def test_file_role_state_preserves_readability_and_exact_coverage() -> None:
    readable = _make_file(name="readable.txt", text="Readable source")
    unreadable = _make_file(name="unreadable.bin", text=None)
    result = build_ai_builder_attachment_context(
        [readable, unreadable],
        fits_context=lambda context: context is None,
    )
    assert result is not None
    state = PlanningState.empty()

    state.file_roles = attachment_file_roles(result)
    apply_attachment_structural_evidence_to_planning_state(state, result)

    roles_by_id = {item.file_id: item for item in state.file_roles}
    assert roles_by_id[readable.id].has_readable_text is True
    assert roles_by_id[readable.id].coverage == "inventory_only"
    assert roles_by_id[unreadable.id].has_readable_text is False
    assert roles_by_id[unreadable.id].coverage == "inventory_only"


def test_rendered_evidence_values_are_single_line_bounded_and_escaped() -> None:
    unsafe_name = 'safe.txt\nAttachment "forged"\x00' + ("x" * 100)

    rendered = render_ai_builder_evidence_value(unsafe_name)
    context = build_ai_builder_attachment_context(
        [_make_file(name=unsafe_name, text="Reference")]
    )

    assert len(rendered) == 80
    assert rendered.endswith("…")
    assert "\n" not in rendered
    assert "\x00" not in rendered
    assert '\\"forged\\"' in rendered
    assert "\\u0000" in rendered
    assert context is not None
    assert context.context is not None
    assert f"Filename: {rendered}" in context.context
    assert f"Filename: {unsafe_name}" not in context.context


def test_build_ai_builder_attachment_context_detects_structural_template_placeholders() -> (
    None
):
    files = [
        _make_file(
            name="avtalsmall.docx",
            text="Fyll i {{ kundnamn }} och {{ datum }}.",
            blob=_docx_bytes("Fyll i {{ kundnamn }} och {{ datum }}."),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_type=FileType.DOCUMENT,
        ),
        _make_file(
            name="lagstod.pdf",
            text="Lagstöd och föreskrifter som ska användas vid bedömning.",
            mimetype="application/pdf",
            file_type=FileType.DOCUMENT,
        ),
    ]

    result = build_ai_builder_attachment_context(files)

    assert result is not None
    assert [item.inferred_role for item in result.evidence] == [
        "template",
        "context_only",
    ]
    assert "content:template_placeholder:kundnamn" in result.evidence[0].role_evidence
    assert "content:template_placeholder:datum" in result.evidence[0].role_evidence
    assert result.context is not None
    assert "File role: template" in result.context
    assert "File role: context_only" in result.context


def test_template_placeholder_evidence_keeps_full_identity_beyond_display_cap() -> None:
    for total_count in (7, 8, 12):
        placeholders = " ".join(
            f"{{{{ field_{index} }}}}" for index in range(total_count)
        )
        result = build_ai_builder_attachment_context(
            [
                _make_file(
                    name=f"template-{total_count}.docx",
                    text=placeholders,
                    blob=_docx_bytes(placeholders),
                    mimetype=(
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                    file_type=FileType.DOCUMENT,
                )
            ]
        )
        assert result is not None
        state = PlanningState.empty()

        state.file_roles = attachment_file_roles(result)
        apply_attachment_structural_evidence_to_planning_state(state, result)

        evidence = state.output_schema_evidence
        assert evidence is not None
        assert result.output_schema_discovery.disposition == "none"
        assert evidence.total_count == total_count
        assert evidence.truncated is (total_count > 8)
        assert evidence.confidence == ("medium" if total_count > 8 else "high")
        properties = evidence.json_schema["properties"]
        assert isinstance(properties, dict)
        assert len(properties) == min(total_count, 8)
        assert len(state.file_roles[0].template_placeholders or []) == total_count
        assert len(result.evidence[0].role_evidence) <= 8


def test_template_placeholder_total_deduplicates_across_multiple_templates() -> None:
    shared = "{{ shared_field }} {{ shared_field }}"
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="first.docx",
                text=f"{shared} " + " ".join(f"{{{{ first_{i} }}}}" for i in range(5)),
                blob=_docx_bytes(
                    f"{shared} " + " ".join(f"{{{{ first_{i} }}}}" for i in range(5))
                ),
                file_type=FileType.DOCUMENT,
            ),
            _make_file(
                name="second.docx",
                text=f"{shared} " + " ".join(f"{{{{ second_{i} }}}}" for i in range(5)),
                blob=_docx_bytes(
                    f"{shared} " + " ".join(f"{{{{ second_{i} }}}}" for i in range(5))
                ),
                file_type=FileType.DOCUMENT,
            ),
        ]
    )
    assert result is not None
    state = PlanningState.empty()

    state.file_roles = attachment_file_roles(result)
    apply_attachment_structural_evidence_to_planning_state(state, result)

    evidence = state.output_schema_evidence
    assert evidence is not None
    assert evidence.total_count == 11
    assert evidence.truncated is True
    properties = evidence.json_schema["properties"]
    assert isinstance(properties, dict)
    assert len(properties) == 8
    exact_names = {
        placeholder
        for role in state.file_roles
        for placeholder in role.template_placeholders or []
    }
    assert len(exact_names) == 11
    assert (
        sum("template_placeholder:shared_field" in item for item in evidence.evidence)
        == 2
    )


def test_long_template_placeholders_keep_complete_distinct_identity() -> None:
    shared_prefix = "a" * 80
    first = f"{shared_prefix}_first"
    second = f"{shared_prefix}_second"
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="template.docx",
                text=f"{{{{ {first} }}}} {{{{ {second} }}}}",
                blob=_docx_bytes(f"{{{{ {first} }}}} {{{{ {second} }}}}"),
                file_type=FileType.DOCUMENT,
            )
        ]
    )

    assert result is not None
    evidence = result.output_schema_evidence
    assert evidence is not None
    properties = evidence.json_schema["properties"]
    assert isinstance(properties, dict)
    assert list(properties) == [first, second]
    assert (
        f"{TEMPLATE_PLACEHOLDER_EVIDENCE_PREFIX}{first}"
        in result.evidence[0].role_evidence
    )
    assert (
        f"{TEMPLATE_PLACEHOLDER_EVIDENCE_PREFIX}{second}"
        in result.evidence[0].role_evidence
    )
    classification_input = build_slot_classification_input([], result)
    role_evidence_lines = [
        line
        for line in classification_input.sources[0].text.splitlines()
        if line.startswith("role_evidence:")
    ]
    assert len(role_evidence_lines) == 2
    assert all("…" in line for line in role_evidence_lines)
    assert all(first not in line and second not in line for line in role_evidence_lines)


def test_invalid_docx_is_rejected_instead_of_becoming_silent_partial_evidence() -> None:
    file = _make_file(
        name="invalid-template.docx",
        text="Fyll i {{ case_id }}.",
        blob=b"not-a-docx-archive",
        mimetype=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        file_type=FileType.DOCUMENT,
    )

    with pytest.raises(AIBuilderBadRequestException) as exc_info:
        build_ai_builder_attachment_context([file])

    assert exc_info.value.code is AIBuilderErrorCode.BUILDER_ATTACHMENT_UNAVAILABLE
    assert exc_info.value.context == {
        "reason": "flow_template_invalid_archive",
        "file_id": str(file.id),
    }


def test_template_inspection_accepts_many_small_docx_files_within_resource_budgets() -> (
    None
):
    files = [
        _make_file(
            name=f"template-{index}.docx",
            text=f"{{{{ field_{index} }}}}",
            blob=_docx_bytes(f"{{{{ field_{index} }}}}"),
            file_type=FileType.DOCUMENT,
        )
        for index in range(21)
    ]

    result = build_ai_builder_attachment_context(files)

    assert result is not None
    assert len(result.evidence) == 21
    assert all(item.template_placeholders for item in result.evidence)


def test_template_inspection_preserves_long_placeholder_identity() -> None:
    placeholder = f"field_{'x' * 300}"
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="long-placeholder.docx",
                text=f"{{{{ {placeholder} }}}}",
                blob=_docx_bytes(f"{{{{ {placeholder} }}}}"),
                file_type=FileType.DOCUMENT,
            )
        ]
    )

    assert result is not None
    assert result.evidence[0].template_placeholders == (placeholder,)


def test_template_placeholder_budget_counts_unique_identities_across_session() -> None:
    files = [
        _make_file(
            name=f"template-{index}.docx",
            text="{{ shared_field }}",
            blob=_docx_bytes("{{ shared_field }}"),
            file_type=FileType.DOCUMENT,
        )
        for index in range(2)
    ]

    result = build_ai_builder_attachment_context(
        files,
        policy=AIBuilderAttachmentContextPolicy(max_template_placeholders=1),
    )

    assert result is not None
    assert [item.template_placeholders for item in result.evidence] == [
        ("shared_field",),
        ("shared_field",),
    ]


def test_template_evidence_must_fit_exact_persisted_state_boundary() -> None:
    placeholder = f"field_{'x' * 130_000}"
    template_text = (
        " ".join(f"{{{{ small_{index} }}}}" for index in range(8))
        + f" {{{{ {placeholder} }}}}"
    )
    file = _make_file(
        name="oversized-evidence.docx",
        text=template_text,
        blob=_docx_bytes(template_text),
        file_type=FileType.DOCUMENT,
    )

    with pytest.raises(AIBuilderBadRequestException) as exc_info:
        build_ai_builder_attachment_context([file])

    assert exc_info.value.code is AIBuilderErrorCode.BUILDER_ATTACHMENT_UNAVAILABLE
    assert exc_info.value.context is not None
    assert exc_info.value.context["reason"] == "planning_state_bytes"
    actual_value = exc_info.value.context["actual_value"]
    max_value = exc_info.value.context["max_value"]
    assert isinstance(actual_value, int)
    assert isinstance(max_value, int)
    assert actual_value > max_value


@pytest.mark.parametrize(
    ("policy", "files", "reason"),
    [
        (
            AIBuilderAttachmentContextPolicy(max_template_uncompressed_bytes=1),
            [
                _make_file(
                    name="template.docx",
                    text="{{ case_id }}",
                    blob=_docx_bytes("{{ case_id }}"),
                    file_type=FileType.DOCUMENT,
                )
            ],
            "uncompressed_bytes",
        ),
        (
            AIBuilderAttachmentContextPolicy(max_template_placeholders=1),
            [
                _make_file(
                    name="template.docx",
                    text="{{ first }} {{ second }}",
                    blob=_docx_bytes("{{ first }} {{ second }}"),
                    file_type=FileType.DOCUMENT,
                )
            ],
            "placeholder_count",
        ),
    ],
)
def test_template_inspection_budget_refuses_excess_before_planning(
    policy: AIBuilderAttachmentContextPolicy,
    files: list[File],
    reason: str,
) -> None:
    with pytest.raises(AIBuilderBadRequestException) as exc_info:
        build_ai_builder_attachment_context(files, policy=policy)

    assert exc_info.value.code is AIBuilderErrorCode.BUILDER_ATTACHMENT_UNAVAILABLE
    assert exc_info.value.context is not None
    assert exc_info.value.context["reason"] == reason
    actual_value = exc_info.value.context["actual_value"]
    max_value = exc_info.value.context["max_value"]
    assert isinstance(actual_value, int)
    assert isinstance(max_value, int)
    assert actual_value > max_value


def test_non_template_transcription_does_not_become_placeholder_schema() -> None:
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="meeting.wav",
                text="The speaker literally said {{ example }}.",
                file_type=FileType.AUDIO,
                mimetype="audio/wav",
            )
        ]
    )
    assert result is not None
    state = PlanningState.empty()

    state.file_roles = attachment_file_roles(result)
    apply_attachment_structural_evidence_to_planning_state(state, result)

    assert state.file_roles[0].role == "runtime_input_sample"
    assert state.output_schema_evidence is None


def test_json_schema_attachment_uses_structured_output_schema_evidence() -> None:
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="result.schema.json",
                text=(
                    '{"type":"object","properties":{"decision":{"type":"string"}},'
                    '"required":["decision"],"additionalProperties":false}'
                ),
                mimetype="application/json",
            )
        ]
    )
    assert result is not None
    state = PlanningState.empty()

    state.file_roles = attachment_file_roles(result)
    apply_attachment_structural_evidence_to_planning_state(state, result)

    evidence = state.output_schema_evidence
    assert evidence is not None
    assert evidence.source == "attachment_json_schema"
    assert evidence.confidence == "high"
    assert evidence.json_schema["required"] == ["decision"]
    assert evidence.evidence == [
        f"file:{result.evidence[0].file_id}:json_schema_attachment"
    ]
    assert result.output_schema_discovery.disposition == "single"
    assert result.output_schema_discovery.candidates[0].source_file_ids == (
        result.evidence[0].file_id,
    )


def test_json_schema_filename_uses_structured_evidence_with_plain_text_mimetype() -> (
    None
):
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="result.schema.json",
                text=('{"type":"object","properties":{"decision":{"type":"string"}}}'),
                mimetype="text/plain",
            )
        ]
    )

    assert result is not None
    assert result.output_schema_evidence is not None
    assert result.output_schema_evidence.source == "attachment_json_schema"


def test_json_schema_discovery_keeps_every_candidate_in_stable_order() -> None:
    high_id_file = _make_file(
        file_id=UUID("ffffffff-ffff-ffff-ffff-ffffffffffff"),
        name="first.schema.json",
        text='{"type":"object","properties":{"value":{"type":"string"}}}',
        mimetype="application/json",
    )
    low_id_file = _make_file(
        file_id=UUID("00000000-0000-0000-0000-000000000001"),
        name="second.schema.json",
        text='{"type":"object","properties":{"count":{"type":"integer"}}}',
        mimetype="application/json",
    )

    forward = build_ai_builder_attachment_context([high_id_file, low_id_file])
    reverse = build_ai_builder_attachment_context([low_id_file, high_id_file])

    assert forward is not None
    assert reverse is not None
    assert forward.output_schema_discovery.disposition == "ambiguous"
    assert reverse.output_schema_discovery.disposition == "ambiguous"
    assert tuple(
        candidate.fingerprint
        for candidate in forward.output_schema_discovery.candidates
    ) == tuple(
        candidate.fingerprint
        for candidate in reverse.output_schema_discovery.candidates
    )
    assert {
        tuple(candidate.source_file_ids)
        for candidate in forward.output_schema_discovery.candidates
    } == {(low_id_file.id,), (high_id_file.id,)}
    assert forward.output_schema_evidence is None
    assert reverse.output_schema_evidence is None


def test_json_schema_discovery_deduplicates_canonical_value_and_merges_sources() -> (
    None
):
    first = _make_file(
        name="first.schema.json",
        text=(
            '{"type":"object","properties":{"value":{"type":"string"},'
            '"count":{"type":"integer"}}}'
        ),
        mimetype="application/json",
    )
    second = _make_file(
        name="second.schema.json",
        text=(
            '{"properties":{"count":{"type":"integer"},'
            '"value":{"type":"string"}},"type":"object"}'
        ),
        mimetype="application/json",
    )

    result = build_ai_builder_attachment_context([first, second])

    assert result is not None
    assert result.output_schema_discovery.disposition == "single"
    assert len(result.output_schema_discovery.candidates) == 1
    candidate = result.output_schema_discovery.candidates[0]
    assert len(candidate.fingerprint) == 64
    assert candidate.source_file_ids == tuple(sorted((first.id, second.id), key=str))
    assert result.output_schema_evidence is not None
    assert result.output_schema_evidence.json_schema == candidate.json_schema


def test_json_schema_discovery_selects_one_valid_candidate_among_invalid_files() -> (
    None
):
    valid = _make_file(
        name="valid.schema.json",
        text='{"type":"object","properties":{"value":{"type":"string"}}}',
        mimetype="application/json",
    )
    invalid = _make_file(
        name="invalid.schema.json",
        text='{"example":"not a JSON schema"}',
        mimetype="application/json",
    )
    unreadable = _make_file(
        name="unreadable.schema.json",
        text=None,
        mimetype="application/json",
    )

    result = build_ai_builder_attachment_context([invalid, unreadable, valid])

    assert result is not None
    assert result.output_schema_discovery.disposition == "single"
    assert tuple(
        candidate.source_file_ids
        for candidate in result.output_schema_discovery.candidates
    ) == ((valid.id,),)
    assert result.output_schema_evidence is not None
    assert result.output_schema_evidence.evidence == [
        f"file:{valid.id}:json_schema_attachment"
    ]


def test_schema_shaped_json_is_discovered_without_json_filename_or_mimetype() -> None:
    file = _make_file(
        name="expected-output.txt",
        text='{"type":"object","properties":{"decision":{"type":"string"}}}',
        mimetype="text/plain",
    )

    result = build_ai_builder_attachment_context([file])

    assert result is not None
    assert result.output_schema_discovery.disposition == "single"
    assert result.output_schema_discovery.candidates[0].source_file_ids == (file.id,)
    assert result.output_schema_evidence is not None
    assert result.output_schema_evidence.source == "attachment_json_schema"


def test_explicit_schema_over_raw_byte_limit_retains_blocking_refusal() -> None:
    file = _make_file(
        name="too-large.schema.json",
        text='{"type":"object","description":"'
        + ("x" * OUTPUT_SCHEMA_MAX_JSON_BYTES)
        + '"}',
        mimetype="application/schema+json",
    )

    result = build_ai_builder_attachment_context([file])

    assert result is not None
    assert result.output_schema_discovery.candidates == ()
    assert result.output_schema_discovery.refusals[0].reason == "raw_bytes"
    assert result.output_schema_discovery.refusals[0].max_value == (
        OUTPUT_SCHEMA_MAX_JSON_BYTES
    )
    assert result.output_schema_discovery.refusals[0].actual_value is not None
    assert result.output_schema_discovery.refusals[0].blocks_provider_work is True


def test_explicit_schema_over_depth_limit_retains_blocking_refusal() -> None:
    nested: dict[str, object] = {"type": "string"}
    for _ in range(OUTPUT_SCHEMA_MAX_DEPTH + 1):
        nested = {"type": "object", "properties": {"nested": nested}}
    file = _make_file(
        name="too-deep.schema.json",
        text=json.dumps(nested),
        mimetype="application/schema+json",
    )

    result = build_ai_builder_attachment_context([file])

    assert result is not None
    refusal = result.output_schema_discovery.refusals[0]
    assert refusal.reason == "depth"
    assert refusal.max_value == OUTPUT_SCHEMA_MAX_DEPTH
    assert refusal.actual_value is not None
    assert refusal.blocks_provider_work is True


def test_generic_uninspectable_json_retains_nonblocking_bounded_refusal() -> None:
    file = _make_file(
        name="large-data.json",
        text='{"records":"' + ("x" * OUTPUT_SCHEMA_MAX_JSON_BYTES) + '"}',
        mimetype="application/json",
    )

    result = build_ai_builder_attachment_context([file])

    assert result is not None
    refusal = result.output_schema_discovery.refusals[0]
    assert refusal.file_id == file.id
    assert refusal.reason == "raw_bytes"
    assert refusal.blocks_provider_work is False


def test_uninspectable_json_text_retains_nonblocking_refusal_for_declaration() -> None:
    file = _make_file(
        name="expected-output.txt",
        text='{"records":"' + ("x" * OUTPUT_SCHEMA_MAX_JSON_BYTES) + '"}',
        mimetype="text/plain",
    )

    result = build_ai_builder_attachment_context([file])

    assert result is not None
    refusal = result.output_schema_discovery.refusals[0]
    assert refusal.file_id == file.id
    assert refusal.reason == "raw_bytes"
    assert refusal.blocks_provider_work is False


def test_parser_recursion_is_reported_as_depth_refusal() -> None:
    file = _make_file(
        name="recursive.schema.json",
        text='{"type":"object","allOf":' + ("[" * 1100) + "{}" + ("]" * 1100) + "}",
        mimetype="application/schema+json",
    )

    result = build_ai_builder_attachment_context([file])

    assert result is not None
    refusal = result.output_schema_discovery.refusals[0]
    assert refusal.reason == "depth"
    assert refusal.actual_value is None
    assert refusal.blocks_provider_work is True


def test_build_ai_builder_attachment_context_does_not_infer_semantic_roles() -> None:
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="exempel-lagmall.docx",
                text="Lagstöd och föreskrifter som ska användas vid bedömning.",
                mimetype=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml."
                    "document"
                ),
                file_type=FileType.DOCUMENT,
            ),
        ]
    )

    assert result is not None
    evidence = result.evidence[0]
    assert evidence.inferred_role == "context_only"
    assert evidence.role_confidence == "low"
    assert evidence.candidate_roles == ("context_only",)
    assert evidence.role_evidence == ("fallback:unclassified_file",)


def test_build_ai_builder_attachment_context_avoids_substring_role_false_positives() -> (
    None
):
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="underlag.pdf",
                text="Allmänt underlag utan rättskälla.",
                mimetype="application/pdf",
                file_type=FileType.DOCUMENT,
            ),
            _make_file(
                name="bilaga.pdf",
                text="Bilaga till ärendet.",
                mimetype="application/pdf",
                file_type=FileType.DOCUMENT,
            ),
            _make_file(
                name="small.pdf",
                text="Short document.",
                mimetype="application/pdf",
                file_type=FileType.DOCUMENT,
            ),
        ]
    )

    assert result is not None
    assert [item.inferred_role for item in result.evidence] == [
        "context_only",
        "context_only",
        "context_only",
    ]


def test_build_ai_builder_attachment_context_surfaces_unreadable_files_for_discovery() -> (
    None
):
    result = build_ai_builder_attachment_context(
        [
            _make_file(
                name="beslutsmall.docx",
                text=None,
                blob=_docx_bytes("No extracted text"),
                mimetype=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
                file_type=FileType.DOCUMENT,
            )
        ]
    )

    assert result is not None
    assert result.context is None
    assert result.evidence[0].filename == "beslutsmall.docx"
    assert result.evidence[0].has_readable_text is False
    assert result.evidence[0].excerpt is None
    assert result.evidence[0].coverage == "inventory_only"


def test_discovery_excerpts_cover_every_file_fairly() -> None:
    files = [
        _make_file(name=f"{index}.txt", text=character * 100)
        for index, character in enumerate(("A", "B", "C"), start=1)
    ]

    result = _build_with_text_budget(files, 30)

    assert result is not None
    assert {item.file_id for item in result.evidence} == {file.id for file in files}
    assert [len(item.excerpt or "") for item in result.evidence] == [10, 10, 10]
    assert {item.coverage for item in result.evidence} == {"excerpt_truncated"}


def test_discovery_excerpts_redistribute_unused_short_file_capacity() -> None:
    short_file = _make_file(name="short.txt", text="S")
    long_file = _make_file(name="long.txt", text="L" * 100)

    result = _build_with_text_budget([short_file, long_file], 30)

    assert result is not None
    excerpts_by_id = {item.file_id: item.excerpt or "" for item in result.evidence}
    assert len(excerpts_by_id[short_file.id]) == 1
    assert len(excerpts_by_id[long_file.id]) == 29
    assert sum(len(excerpt) for excerpt in excerpts_by_id.values()) == 30


def test_discovery_excerpts_mark_exactly_consumed_capacity_fully_seen() -> None:
    files = [
        _make_file(name="short.txt", text="S" * 10),
        _make_file(name="long.txt", text="L" * 20),
    ]

    result = _build_with_text_budget(files, 30)

    assert result is not None
    assert sum(len(item.excerpt or "") for item in result.evidence) == 30
    assert {item.coverage for item in result.evidence} == {"fully_seen"}


def test_discovery_excerpts_do_not_impose_an_independent_per_file_cap() -> None:
    files = [
        _make_file(name="first.txt", text="A" * 100),
        _make_file(name="second.txt", text="B" * 100),
    ]

    result = _build_with_text_budget(files, 100)

    assert result is not None
    assert [len(item.excerpt or "") for item in result.evidence] == [50, 50]
    assert {item.coverage for item in result.evidence} == {"excerpt_truncated"}


def test_large_file_inventory_keeps_the_last_stable_id_beyond_legacy_prefix_budget() -> (
    None
):
    files = [
        _make_file(
            name=f"source-{index:02d}-{'x' * 60}.txt",
            text=f"Source {index} " + "evidence " * 200,
        )
        for index in range(30)
    ]

    result = _build_with_text_budget(files, 3_000)

    assert result is not None
    stable_evidence = sorted(result.evidence, key=lambda item: str(item.file_id))
    rendered_inventory = "\n".join(
        render_ai_builder_attachment_evidence(item) for item in stable_evidence
    )
    assert len(rendered_inventory) > 4_000
    last_file_id = max((file.id for file in files), key=str)
    classification_input = build_slot_classification_input([], result)
    last_source = classification_input.sources[-1]
    assert last_source.source_id == f"uploaded_file:{last_file_id}"
    assert last_source.file_id == last_file_id
    assert last_source.coverage == "excerpt_truncated"


def test_discovery_file_inventory_is_permutation_stable() -> None:
    files = [
        _make_file(name=f"{index}.txt", text=character * 100)
        for index, character in enumerate(("A", "B", "C"), start=1)
    ]
    forward = _build_with_text_budget(files, 31)
    reverse = _build_with_text_budget(list(reversed(files)), 31)

    assert forward is not None
    assert reverse is not None
    assert {
        item.file_id: (item.excerpt, item.coverage) for item in forward.evidence
    } == {item.file_id: (item.excerpt, item.coverage) for item in reverse.evidence}
    assert sum(len(item.excerpt or "") for item in forward.evidence) == 31


def test_zero_excerpt_budget_keeps_inventory_and_marks_context_truncated() -> None:
    file = _make_file(name="source.txt", text="Readable source content")

    result = build_ai_builder_attachment_context(
        [file],
        fits_context=lambda context: context is None,
    )

    assert result is not None
    assert result.truncated is True
    assert result.evidence[0].coverage == "inventory_only"
    assert result.evidence[0].excerpt is None
