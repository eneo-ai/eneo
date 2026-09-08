"""Build DOCX templates with content controls for tests, and read them back.

python-docx's ``document.paragraphs`` skips paragraphs inside structured
document tags, so readers here walk the body XML directly.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from eneo.flows.runtime.document_rendering.docx_content_controls import (
    append_rich_control,
    append_text_control,
)

_EXPRESSION = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def control_template_bytes(
    *,
    rich: Iterable[str | tuple[str, str, str]] = (),
    text: Iterable[str | tuple[str, str, str]] = (),
    heading_level: int | None = 1,
) -> bytes:
    """A template with inline text controls and body-level rich controls.

    Each entry is a tag or ``(tag, label, hint)``. Text controls sit in a
    "Label: " paragraph; rich controls follow a heading with their label
    when ``heading_level`` is set.
    """

    document = Document()
    for entry in text:
        tag, label, hint = _entry(entry)
        paragraph = document.add_paragraph(f"{label}: ")
        append_text_control(paragraph, tag=tag, label=label, hint=hint)
    for entry in rich:
        tag, label, hint = _entry(entry)
        if heading_level is not None:
            document.add_heading(label, level=heading_level)
        append_rich_control(document, tag=tag, label=label, hint=hint)
    return _bytes(document)


def template_bytes_from_expressions(text: str) -> bytes:
    """A template with one inline text control per distinct ``{{ name }}`` in ``text``.

    Lets fixtures keep describing a template as prose with expressions; the
    controls are what the runtime reads.
    """

    document = Document()
    prose = _EXPRESSION.sub("", text).strip()
    paragraph = document.add_paragraph(f"{prose} " if prose else "")
    seen: list[str] = []
    for name in _EXPRESSION.findall(text):
        if name in seen:
            continue
        seen.append(name)
        append_text_control(paragraph, tag=name, label=name, hint=name)
    return _bytes(document)


def docx_paragraphs(blob_or_document: bytes | object) -> list[Paragraph]:
    """Every body paragraph in reading order, including those inside controls."""

    document = _document(blob_or_document)
    return [
        Paragraph(element, document._body)
        for element in document.element.body.iter(qn("w:p"))
        if element.getparent().tag != qn("w:tc")
    ]


def docx_paragraph_texts(blob_or_document: bytes | object) -> list[str]:
    """Paragraph texts read from the XML, so runs inside inline controls count."""

    return [
        _paragraph_text(paragraph) for paragraph in docx_paragraphs(blob_or_document)
    ]


def docx_styled_paragraphs(blob_or_document: bytes | object) -> list[tuple[str, str]]:
    """``(style name, text)`` per non-empty body paragraph, controls included."""

    return [
        (paragraph.style.name, _paragraph_text(paragraph))
        for paragraph in docx_paragraphs(blob_or_document)
        if _paragraph_text(paragraph)
    ]


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(text.text or "" for text in paragraph._p.iter(qn("w:t")))


def docx_tables(blob_or_document: bytes | object) -> list[Table]:
    document = _document(blob_or_document)
    return [
        Table(element, document._body)
        for element in document.element.body.iter(qn("w:tbl"))
    ]


def _entry(entry: str | tuple[str, str, str]) -> tuple[str, str, str]:
    if isinstance(entry, str):
        return entry, entry, f"Fyll i {entry}"
    return entry


def _document(blob_or_document: bytes | object):
    if isinstance(blob_or_document, bytes):
        return Document(io.BytesIO(blob_or_document))
    return blob_or_document


def _bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
